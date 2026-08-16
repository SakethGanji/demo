"""Discovery contract fixes — literal search, real pages, honest usage, item GETs.

Four independent defects, all in the discovery surface a catalog UI drives:

1. ``GET /search/columns`` concatenated the raw query into an ILIKE pattern, so
   a user typing ``%`` or ``_`` got somebody else's answer.
2. The two sheet-metadata list routes fabricated the ``Page`` envelope
   (``limit=len(items), offset=0``) and dropped ``?limit=``/``?offset=``.
3. ``GET /datasets/{id}/usage`` counted denied and failed requests as usage.
4. There was no ``GET`` on the sheet-metadata or column-dictionary ITEM paths,
   so reading one record meant listing everything and filtering client-side.
"""

from __future__ import annotations

import json

from openpyxl import Workbook

from conftest import (
    DEFAULT_TEAM_ID, XLSX_MIME, auth, create_team_user, rid, upload_file,
    upload_inline,
)


async def _upload_wb(client, admin_id, path, *, dataset_id=None):
    return await upload_file(client, admin_id, path, name="book.xlsx",
                             content_type=XLSX_MIME, dataset_id=dataset_id)


# ---------------------------------------------------------------------------
# 1. Column search matches the query literally
# ---------------------------------------------------------------------------


async def test_column_search_treats_percent_and_underscore_as_literal_text(
        client, admin_id):
    """A search box must not be a pattern language the user cannot see.

    The query was concatenated into ``'%' || :q || '%'``, so Postgres read the
    user's own ``_`` as "any character" and their ``%`` as "anything at all".
    Searching ``qty_eur`` quietly returned ``qtyxeur`` too, and a lone ``%``
    returned the caller's entire visible catalog dressed up as search hits —
    the kind of wrong answer nobody can spot, because the results look real.
    """
    h = auth(admin_id)
    marker = f"c{rid()}"
    exact, decoy = f"{marker}_qty", f"{marker}xqty"
    ds = (await upload_inline(client, admin_id, json.dumps(
        [{exact: 1, decoy: 2}])))["dataset_id"]

    async def search(q):
        r = await client.get("/api/v1/search/columns", params={"q": q}, headers=h)
        assert r.status_code == 200, r.text
        return [hit["normalized_name"] for hit in r.json()["items"]
                if hit["dataset_id"] == ds]

    # `_` is a literal underscore, not "any single character".
    assert await search(exact) == [exact]
    # ...and the decoy is genuinely there, so the assertion above has teeth.
    assert await search(decoy) == [decoy]
    assert sorted(await search(marker)) == [exact, decoy]

    # `%` matches nothing, because no column name contains one.
    assert await search("%") == []
    assert await search(f"%{marker}") == []
    # A backslash is data too — escaping it must not smuggle a wildcard back in.
    assert await search(f"\\{marker}") == []


# ---------------------------------------------------------------------------
# 2. The sheet-metadata lists are real pages
# ---------------------------------------------------------------------------


def _wb_with_sheets(path, names):
    wb = Workbook()
    for i, name in enumerate(names):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        ws.append(["Item", "Cost"])
        ws.append(["rent", 50])
    wb.save(path)


async def test_sheet_metadata_list_honours_limit_and_offset(client, admin_id, tmp_path):
    """The ``Page`` envelope is a promise: ``total`` is the whole collection and
    ``limit``/``offset`` are what the client asked for.

    This route ignored both query parameters and echoed ``limit=len(items)``,
    so a UI paging through a dataset's sheets silently got everything on page
    one and computed "1 page of N" from a limit the server invented. A dataset
    with no recorded metadata reported ``limit=0`` — a page size the service's
    own pagination contract (``ge=1``) forbids, and a division by zero waiting
    to happen in any client that derives a page count from it.
    """
    h = auth(admin_id)
    p = tmp_path / "book.xlsx"
    _wb_with_sheets(p, ["Alpha", "Bravo", "Charlie"])
    ds = (await _upload_wb(client, admin_id, p))["dataset_id"]
    url = f"/api/v1/datasets/{ds}/sheet-metadata"

    # Empty collection: a page size the client could actually send.
    empty = (await client.get(url, headers=h)).json()
    assert empty["items"] == [] and empty["total"] == 0
    assert empty["limit"] == 50 and empty["offset"] == 0

    for sheet in ("Alpha", "Bravo", "Charlie"):
        r = await client.put(f"{url}/{sheet}", headers=h,
                             json={"grain": f"one row per {sheet.lower()}"})
        assert r.status_code == 200, r.text

    body = (await client.get(url, params={"limit": 2}, headers=h)).json()
    assert [m["sheet_key"] for m in body["items"]] == ["alpha", "bravo"]
    assert body["total"] == 3 and body["limit"] == 2 and body["offset"] == 0

    body = (await client.get(url, params={"limit": 2, "offset": 1}, headers=h)).json()
    assert [m["sheet_key"] for m in body["items"]] == ["bravo", "charlie"]
    assert body["total"] == 3 and body["limit"] == 2 and body["offset"] == 1

    # Past the end: still a well-formed page, not a lie about the total.
    body = (await client.get(url, params={"limit": 2, "offset": 9}, headers=h)).json()
    assert body["items"] == [] and body["total"] == 3
    assert body["limit"] == 2 and body["offset"] == 9


async def test_column_dictionary_list_honours_limit_and_offset(client, admin_id):
    """Same envelope promise for the data dictionary, which is the one that
    actually gets long — a wide sheet documents hundreds of columns, and the
    route returned every one of them however small a page was requested."""
    h = auth(admin_id)
    ds = (await upload_inline(client, admin_id, json.dumps(
        [{"amount": 1, "book": "b", "currency": "EUR", "region": "EU"}])))["dataset_id"]
    base = f"/api/v1/datasets/{ds}/sheet-metadata/data/columns"

    empty = (await client.get(base, headers=h)).json()
    assert empty["total"] == 0 and empty["limit"] == 50 and empty["offset"] == 0

    for col in ("amount", "book", "currency", "region"):
        r = await client.put(f"{base}/{col}", headers=h,
                             json={"business_name": col.title()})
        assert r.status_code == 200, r.text

    body = (await client.get(base, params={"limit": 2, "offset": 1}, headers=h)).json()
    assert [e["column_name"] for e in body["items"]] == ["book", "currency"]
    assert body["total"] == 4 and body["limit"] == 2 and body["offset"] == 1


# ---------------------------------------------------------------------------
# 3. Usage counts successful activity only
# ---------------------------------------------------------------------------


async def test_dataset_usage_excludes_denied_and_failed_requests(client, admin_id):
    """"How much is this dataset used?" must not be answerable by being refused.

    The audit middleware records the FINAL response of every mutating or
    download request, including 403s, 404s and 422s. Usage counted those rows
    unfiltered, so a viewer repeatedly bounced off a write, or a bot probing a
    dataset it cannot see, showed up as real writes and downloads — and
    ``last_activity_at`` advanced on activity that never happened. That number
    is what a UI (and the MCP context tool) uses to argue a dataset matters.
    """
    h = auth(admin_id)
    ds = (await upload_inline(client, admin_id, '[{"a": 1}]'))["dataset_id"]

    async def usage():
        r = await client.get(f"/api/v1/datasets/{ds}/usage", headers=h)
        assert r.status_code == 200, r.text
        return r.json()

    # One real write and one real download.
    assert (await client.put(f"/api/v1/datasets/{ds}/sheet-metadata/data",
                             headers=h, json={"grain": "one row"})).status_code == 200
    assert (await client.get(f"/api/v1/datasets/{ds}/download",
                             params={"format": "csv", "sheet": "data"},
                             headers=h)).status_code == 200
    before = await usage()
    assert before["downloads"] == 1 and before["writes"] == 1
    assert before["total_events"] == 2 and before["last_activity_at"]

    # Now a denied write, a hidden-dataset download, and a rejected body.
    viewer, _ = await create_team_user(client, admin_id, "viewer",
                                       team_id=DEFAULT_TEAM_ID)
    outsider, _ = await create_team_user(client, admin_id, "editor")
    assert (await client.put(f"/api/v1/datasets/{ds}/sheet-metadata/data",
                             headers=auth(viewer),
                             json={"grain": "x"})).status_code == 403
    assert (await client.get(f"/api/v1/datasets/{ds}/download",
                             params={"format": "csv", "sheet": "data"},
                             headers=auth(outsider))).status_code == 404
    assert (await client.patch(f"/api/v1/datasets/{ds}", headers=h,
                               json={"name": None})).status_code == 422

    # Those failures ARE in the audit trail — the timeline shows them — so the
    # counts below are filtered, not merely missing their rows.
    events = (await client.get(f"/api/v1/datasets/{ds}/timeline",
                               headers=h)).json()["items"]
    codes = {e["details"].get("status_code") for e in events
             if e["event_type"] == "audit"}
    assert 403 in codes and 422 in codes

    after = await usage()
    assert after["downloads"] == 1 and after["writes"] == 1
    assert after["total_events"] == 2
    assert after["last_activity_at"] == before["last_activity_at"]


# ---------------------------------------------------------------------------
# 4. Single records are readable at their own URL
# ---------------------------------------------------------------------------


async def test_one_sheet_metadata_record_is_readable_at_its_own_url(
        client, admin_id, tmp_path):
    """PUT and PATCH addressed this path but GET returned 405.

    Reading back one record meant fetching every sheet's metadata and filtering
    client-side — which also silently breaks once that list is paged. GET here
    normalizes ``sheet_key`` exactly as PUT does, so the URL a client wrote to
    is the URL it can read from.
    """
    h = auth(admin_id)
    p = tmp_path / "book.xlsx"
    _wb_with_sheets(p, ["Alpha", "Bravo"])
    ds = (await _upload_wb(client, admin_id, p))["dataset_id"]
    base = f"/api/v1/datasets/{ds}/sheet-metadata"

    assert (await client.put(f"{base}/Alpha", headers=h, json={
        "grain": "one row per alpha",
        "primary_key_columns": ["item"]})).status_code == 200

    for spelling in ("Alpha", "alpha"):
        r = await client.get(f"{base}/{spelling}", headers=h)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["sheet_key"] == "alpha"
        assert body["grain"] == "one row per alpha"
        assert body["primary_key_columns"] == ["item"]

    # A sheet with no record 404s rather than inventing an empty one.
    r = await client.get(f"{base}/Bravo", headers=h)
    assert r.status_code == 404, r.text
    assert r.json()["code"] == "not_found"

    # Cross-team readers get the same 404 the rest of the dataset gives them.
    outsider, _ = await create_team_user(client, admin_id, "editor")
    assert (await client.get(f"{base}/Alpha",
                             headers=auth(outsider))).status_code == 404


async def test_one_dictionary_entry_is_readable_at_its_own_url(
        client, admin_id, tmp_path):
    """The item GET the DELETE docstring already claimed existed.

    It normalizes the column name like every other verb on this path, so
    ``Business Name`` reads back the entry ``Business Name`` created, and it
    tolerates a column that has since left the schema — documentation outlives
    the column it describes, and a UI showing a stale entry must be able to
    open it.
    """
    h = auth(admin_id)
    def accounts(path, headers):
        wb = Workbook()
        ws = wb.active
        ws.title = "Accounts"
        ws.append(headers)
        ws.append(["Acme", "EU"][: len(headers)])
        wb.save(path)

    p = tmp_path / "accounts.xlsx"
    accounts(p, ["Business Name", "Region"])
    ds = (await _upload_wb(client, admin_id, p))["dataset_id"]
    base = f"/api/v1/datasets/{ds}/sheet-metadata/Accounts/columns"

    assert (await client.put(f"{base}/Business Name", headers=h, json={
        "business_name": "Legal entity", "sensitivity": "internal",
        "allowed_values": ["Acme"]})).status_code == 200

    for spelling in ("Business Name", "business_name"):
        r = await client.get(f"{base}/{spelling}", headers=h)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["column_name"] == "business_name"
        assert body["sheet_key"] == "accounts"
        assert body["business_name"] == "Legal entity"
        assert body["sensitivity"] == "internal"
        assert body["allowed_values"] == ["Acme"]

    # Undocumented column, unknown column, unknown sheet.
    assert (await client.get(f"{base}/Region", headers=h)).status_code == 404
    assert (await client.get(f"{base}/nope", headers=h)).status_code == 404
    assert (await client.get(
        f"/api/v1/datasets/{ds}/sheet-metadata/nope/columns/region",
        headers=h)).status_code == 404

    # In-team viewers can read the dictionary; outsiders cannot see it exists.
    viewer, _ = await create_team_user(client, admin_id, "viewer",
                                       team_id=DEFAULT_TEAM_ID)
    outsider, _ = await create_team_user(client, admin_id, "editor")
    assert (await client.get(f"{base}/business_name",
                             headers=auth(viewer))).status_code == 200
    assert (await client.get(f"{base}/business_name",
                             headers=auth(outsider))).status_code == 404

    # The column leaves the schema; its documentation is still readable, the
    # same way DELETE can still reach it.
    v2 = tmp_path / "accounts_v2.xlsx"
    accounts(v2, ["Region"])
    await _upload_wb(client, admin_id, v2, dataset_id=ds)
    r = await client.get(f"{base}/business_name", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["business_name"] == "Legal entity"
