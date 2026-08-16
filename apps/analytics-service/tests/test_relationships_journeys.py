"""Relationship + join-builder journeys — driven exactly as a UI would (§22–§23).

Every test here is one screen session: the calls appear in the order a browser
makes them, and each step asserts something that is only true BECAUSE of the
step before it. Ids are never fabricated — each one is read out of an earlier
response, which is the same constraint the frontend lives under.

The screens covered:

* **Review inbox** — seed from quality rules, run discovery (background AND
  in-request), page the inbox, confirm one edge and reject another, and prove a
  re-run cannot overturn either verdict.
* **Declare-relationship modal** — every way the form can be wrong, and what
  the second submit of the same pair does.
* **Sheet-rename repair** — a renamed sheet breaks the join, confirm-rename
  repairs it, and the SAME relationship id drives the join again.
* **Join wizard** — pick the confirmed edge, compare inner vs left, narrow the
  projection, execute, read the result back, and find the run in the library.
* **Version pinning** — the wizard's default stops working as the dataset moves
  on, and the two distinct failures a UI must render differently.
* **Publish** — a join output published as a new VERSION of the left dataset.
* **Delete** — the edge behind a saved join is refused until the library row
  that depends on it is cleared.
* **Permission matrix** — viewer vs editor vs outsider on every route.
"""

from __future__ import annotations

from openpyxl import Workbook

from conftest import (
    DEFAULT_TEAM_ID,
    XLSX_MIME,
    auth,
    create_team_user,
    upload_file,
)

PROBLEM = "application/problem+json"


# ---------------------------------------------------------------------------
# Fixture workbook: Customers 1:N Orders, both referencing Regions.
#
# Three sheets so the version has real relationship structure AND multi-sheet
# selection contracts to trip over. The region_code columns are deliberately
# non-unique on Customers/Orders so discovery's uniqueness floor keeps them
# from proposing a bogus Customers<->Orders edge.
# ---------------------------------------------------------------------------


def _sales_workbook(path, *, orders_sheet="Orders", orders_key=True,
                    include_orders=True, orphan_order=False):
    wb = Workbook()
    cust = wb.active
    cust.title = "Customers"
    cust.append(["customer_id", "region_code", "tier"])
    for row in ([1, "EU", "gold"], [2, "US", "silver"], [3, "EU", "gold"]):
        cust.append(row)

    if include_orders:
        orders = wb.create_sheet(orders_sheet)
        header = ["order_id", "customer_id", "region_code", "total"]
        rows = [[10, 1, "EU", 100.0], [11, 2, "US", 40.0], [12, 1, "EU", 60.0]]
        if orphan_order:
            rows.append([13, 999, "US", 5.0])
        if not orders_key:
            header = [c for c in header if c != "customer_id"]
            rows = [[r[0], r[2], r[3]] for r in rows]
        orders.append(header)
        for row in rows:
            orders.append(row)

    regions = wb.create_sheet("Regions")
    regions.append(["region_code", "region_name"])
    for row in (["EU", "Europe"], ["US", "United States"], ["APAC", "Asia Pacific"]):
        regions.append(row)
    wb.save(path)


async def _sales_dataset(client, user_id, tmp_path, name="sales.xlsx", *,
                         dataset_id=None, team_id=DEFAULT_TEAM_ID, **kw):
    path = tmp_path / name
    _sales_workbook(path, **kw)
    body = await upload_file(client, user_id, path, name="sales.xlsx",
                             content_type=XLSX_MIME, dataset_id=dataset_id,
                             team_id=team_id)
    return body["dataset_id"]


async def _declare(client, user_id, dataset_id, *, from_sheet="Orders",
                   from_column="customer_id", to_sheet="Customers",
                   to_column="customer_id", to_dataset_id=None, confirmed=True):
    """The declare-relationship modal's happy path, as a helper."""
    payload = {"from_sheet": from_sheet, "from_column": from_column,
               "to_sheet": to_sheet, "to_column": to_column,
               "confirmed": confirmed}
    if to_dataset_id:
        payload["to_dataset_id"] = to_dataset_id
    r = await client.post(f"/api/v1/datasets/{dataset_id}/relationships",
                          headers=auth(user_id), json=payload)
    assert r.status_code == 201, r.text
    return r.json()


# ---------------------------------------------------------------------------
# 1. The review inbox
# ---------------------------------------------------------------------------


async def test_a_steward_seeds_from_rules_runs_discovery_and_works_the_inbox_to_empty(
        client, admin_id, tmp_path):
    """SCREEN: Relationships > Review inbox.

    Mount lists the inbox; "Seed from rules" and "Run discovery" are the two
    buttons that fill it; each row has Confirm/Reject. The UI would break if
    the inbox stopped paging, if `total` disagreed with the filter chips, or if
    re-running discovery could overturn a verdict a steward already gave.
    """
    from app.shared import worker

    h = auth(admin_id)
    ds = await _sales_dataset(client, admin_id, tmp_path)

    # ---- 1. A steward declares the FK in the quality screen first ----------
    rule = await client.post(f"/api/v1/datasets/{ds}/rules", headers=h, json={
        "name": "orders-customer-fk", "rule_type": "foreign_key",
        "sheet_selector": "Orders", "column_selector": "customer_id",
        "parameters": {"ref_sheet": "Customers", "ref_column": "customer_id"}})
    assert rule.status_code == 201, rule.text
    rule = rule.json()

    # ---- 2. "Seed from rules" — the inbox gains exactly that one edge ------
    seeded = await client.post(f"/api/v1/datasets/{ds}/relationships/seed", headers=h)
    assert seeded.status_code == 200, seeded.text
    assert seeded.json()["created"] == 1
    fk_edge = seeded.json()["relationships"][0]
    assert (fk_edge["from_sheet"], fk_edge["from_column"]) == ("orders", "customer_id")
    assert (fk_edge["to_sheet"], fk_edge["to_column"]) == ("customers", "customer_id")
    # The badge the row renders, and the "why" panel behind it.
    assert fk_edge["method"] == "fk_rule" and fk_edge["status"] == "suggested"
    assert fk_edge["evidence"]["rule_name"] == rule["name"]
    assert fk_edge["confidence"] == 1.0

    inbox = await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h)
    assert inbox.status_code == 200, inbox.text
    assert inbox.json()["total"] == 1
    assert [e["id"] for e in inbox.json()["items"]] == [fk_edge["id"]]

    # ---- 3. "Run discovery" in the background: the response is a handle ----
    queued = await client.post(
        f"/api/v1/datasets/{ds}/relationships/suggest", params={"sync": "false"},
        headers=h)
    assert queued.status_code == 200, queued.text
    job_id = queued.json()["job_id"]
    assert job_id, "a background run must hand back something to poll"
    assert queued.json()["suggested"] == 0            # nothing has run yet
    # The inbox the response echoes is still only the seeded row.
    assert [e["id"] for e in queued.json()["relationships"]] == [fk_edge["id"]]

    progress = await client.get(f"/api/v1/jobs/{job_id}", headers=h)
    assert progress.status_code == 200, progress.text
    assert progress.json()["status"] == "pending"
    assert progress.json()["job_type"] == "relationship_discovery"
    assert progress.json()["dataset_id"] == ds

    assert await worker.run_pending_jobs_once() >= 1

    done = await client.get(f"/api/v1/jobs/{job_id}", headers=h)
    assert done.json()["status"] == "completed"
    assert done.json()["result"]["suggested"] >= 1

    # ---- 4. The inbox now pages — three edges, most confident first --------
    page1 = await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h,
                             params={"status": "suggested", "limit": 2, "offset": 0})
    assert page1.status_code == 200, page1.text
    body = page1.json()
    assert body["total"] == 3 and body["limit"] == 2 and body["offset"] == 0
    assert len(body["items"]) == 2
    confidences = [e["confidence"] for e in body["items"]]
    assert confidences == sorted(confidences, reverse=True)

    page2 = await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h,
                             params={"status": "suggested", "limit": 2, "offset": 2})
    assert page2.json()["total"] == 3 and len(page2.json()["items"]) == 1
    all_ids = [e["id"] for e in body["items"]] + [e["id"] for e in page2.json()["items"]]
    assert len(set(all_ids)) == 3, "paging must not repeat or drop a row"

    edges = {(e["from_sheet"], e["from_column"], e["to_sheet"], e["to_column"]): e
             for e in body["items"] + page2.json()["items"]}
    assert ("orders", "customer_id", "customers", "customer_id") in edges
    assert ("orders", "region_code", "regions", "region_code") in edges
    assert ("customers", "region_code", "regions", "region_code") in edges
    region_edge = edges[("orders", "region_code", "regions", "region_code")]
    unreviewed = edges[("customers", "region_code", "regions", "region_code")]

    # ---- 5. Discovery re-derived the SEEDED pair — same row, same badge ----
    detail = await client.get(
        f"/api/v1/datasets/{ds}/relationships/{fk_edge['id']}", headers=h)
    assert detail.status_code == 200, detail.text
    same = detail.json()
    assert same["id"] == fk_edge["id"] and same["status"] == "suggested"
    # The row a reviewer opened as "from a quality rule" still reads that way:
    # a measurement corroborates a declaration, it does not replace it. The
    # badge, the confidence the inbox sorts on, and the link back to the rule
    # all survive a discovery run.
    assert same["method"] == "fk_rule"
    assert same["confidence"] == 1.0
    assert same["evidence"]["rule_name"] == "orders-customer-fk"
    # ...and the fresh measurement is filed underneath it, not thrown away.
    assert same["evidence"]["statistical"]["coverage"] == 1.0
    assert same["evidence"]["statistical"]["target_uniqueness"] == 1.0
    assert same["evidence"]["statistical"]["name_score"] > 0

    # ---- 6. Review: confirm one, reject another ---------------------------
    confirmed = await client.post(
        f"/api/v1/datasets/{ds}/relationships/{fk_edge['id']}/confirm", headers=h)
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "confirmed"
    assert confirmed.json()["reviewed_by"] == admin_id

    rejected = await client.post(
        f"/api/v1/datasets/{ds}/relationships/{region_edge['id']}/reject", headers=h)
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["status"] == "rejected"

    # ---- 7. The filter chips now add up to the whole inbox ----------------
    counts = {}
    for status in ("suggested", "confirmed", "rejected"):
        r = await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h,
                             params={"status": status})
        counts[status] = r.json()["total"]
    assert counts == {"suggested": 1, "confirmed": 1, "rejected": 1}
    assert sum(counts.values()) == (
        await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h)
    ).json()["total"]

    # ---- 8. Re-running discovery must not overturn either verdict ---------
    again = await client.post(f"/api/v1/datasets/{ds}/relationships/suggest",
                              headers=h)
    assert again.status_code == 200, again.text
    assert again.json()["job_id"]
    assert again.json()["skipped"] == 0        # a small workbook is never capped
    assert again.json()["pairs_examined"] >= 3
    # The response's own inbox holds only the still-unreviewed edge.
    assert [e["id"] for e in again.json()["relationships"]] == [unreviewed["id"]]

    survived = {}
    for edge_id in (fk_edge["id"], region_edge["id"], unreviewed["id"]):
        r = await client.get(f"/api/v1/datasets/{ds}/relationships/{edge_id}",
                             headers=h)
        survived[edge_id] = r.json()["status"]
    assert survived == {fk_edge["id"]: "confirmed",
                        region_edge["id"]: "rejected",
                        unreviewed["id"]: "suggested"}

    # ---- 9. The Confirm button on an already-confirmed row --------------
    r = await client.post(
        f"/api/v1/datasets/{ds}/relationships/{fk_edge['id']}/confirm", headers=h)
    assert r.status_code == 409, r.text
    assert r.headers["content-type"].startswith(PROBLEM)
    assert r.json()["code"] == "invalid-relationship-transition"
    assert r.json()["current_status"] == "confirmed"
    assert r.json()["target_status"] == "confirmed"

    # ...but changing your mind about a rejection is legal, and re-review
    # re-stamps the reviewer.
    r = await client.post(
        f"/api/v1/datasets/{ds}/relationships/{region_edge['id']}/confirm", headers=h)
    assert r.status_code == 200 and r.json()["status"] == "confirmed"


# ---------------------------------------------------------------------------
# 2. The declare-relationship modal
# ---------------------------------------------------------------------------


async def test_the_declare_relationship_modal_handles_every_way_the_form_is_wrong(
        client, admin_id, tmp_path):
    """SCREEN: Relationships > "Declare relationship" modal.

    Four fields, each of which the API validates against the CURRENT version.
    A UI has to render a sheet picker, a column picker and a target-dataset
    picker; every problem response here is what populates or corrects one of
    them. If the codes regressed the modal could only show raw prose.
    """
    h = auth(admin_id)
    ds = await _sales_dataset(client, admin_id, tmp_path)
    base = f"/api/v1/datasets/{ds}/relationships"

    # ---- 1. Submit with no sheet chosen: the API names the choices --------
    r = await client.post(base, headers=h, json={"from_column": "customer_id",
                                                 "to_column": "customer_id"})
    assert r.status_code == 400, r.text
    assert r.headers["content-type"].startswith(PROBLEM)
    assert r.json()["code"] == "sheet-selection-required"
    assert set(r.json()["sheets"]) == {"Customers", "Orders", "Regions"}

    # ---- 2. A sheet that isn't in this version ---------------------------
    r = await client.post(base, headers=h, json={
        "from_sheet": "Invoices", "from_column": "customer_id",
        "to_sheet": "Customers", "to_column": "customer_id"})
    assert r.status_code == 404, r.text

    # ---- 3. A column that isn't on the chosen sheet ----------------------
    r = await client.post(base, headers=h, json={
        "from_sheet": "Orders", "from_column": "client_id",
        "to_sheet": "Customers", "to_column": "customer_id"})
    assert r.status_code == 400, r.text
    assert r.json()["code"] == "unknown-column"
    assert r.json()["column"] == "client_id"
    # The dropdown the modal should have offered in the first place.
    assert set(r.json()["available"]) == {"order_id", "customer_id",
                                          "region_code", "total"}

    # ...and the same check applies to the TARGET side, so a UI can highlight
    # the right field.
    r = await client.post(base, headers=h, json={
        "from_sheet": "Orders", "from_column": "customer_id",
        "to_sheet": "Customers", "to_column": "client_id"})
    assert r.status_code == 400 and r.json()["code"] == "unknown-column"
    assert set(r.json()["available"]) == {"customer_id", "region_code", "tier"}

    # ---- 4. A target dataset the caller cannot read is simply absent -----
    _outsider, other_team = await create_team_user(client, admin_id, "admin")
    foreign = await _sales_dataset(client, admin_id, tmp_path, name="foreign.xlsx",
                                   team_id=other_team)
    editor, _ = await create_team_user(client, admin_id, "editor",
                                       team_id=DEFAULT_TEAM_ID)
    r = await client.post(base, headers=auth(editor), json={
        "from_sheet": "Orders", "from_column": "customer_id",
        "to_dataset_id": foreign, "to_sheet": "Customers",
        "to_column": "customer_id"})
    assert r.status_code == 404, r.text
    assert foreign not in r.text or "not found" in r.text.lower()

    # ---- 5. The correct submission, deliberately left unconfirmed --------
    r = await client.post(base, headers=h, json={
        "from_sheet": "Orders", "from_column": "customer_id",
        "to_sheet": "Customers", "to_column": "customer_id",
        "confirmed": False})
    assert r.status_code == 201, r.text
    edge = r.json()
    assert edge["method"] == "manual" and edge["status"] == "suggested"
    assert edge["to_dataset_id"] == ds        # within-workbook edges self-reference
    assert edge["evidence"]["declared_by"] == admin_id
    assert edge["created_by"] == admin_id and edge["reviewed_by"] is None

    # ---- 6. Submitting the SAME pair again is an update, not a duplicate --
    r = await client.post(base, headers=h, json={
        "from_sheet": "Orders", "from_column": "customer_id",
        "to_sheet": "Customers", "to_column": "customer_id",
        "confirmed": True})
    assert r.status_code == 201, r.text
    assert r.json()["id"] == edge["id"], "the modal must not create a second row"
    assert r.json()["status"] == "confirmed"
    listing = await client.get(base, headers=h)
    assert listing.json()["total"] == 1

    # ---- 7. Re-declaring a REJECTED pair returns it still rejected -------
    # A human verdict outranks a declaration, which is the right call — but the
    # response is a 201 carrying status='rejected', so the modal must read the
    # body rather than trusting the status code.
    await client.post(f"{base}/{edge['id']}/reject", headers=h)
    r = await client.post(base, headers=h, json={
        "from_sheet": "Orders", "from_column": "customer_id",
        "to_sheet": "Customers", "to_column": "customer_id",
        "confirmed": True})
    assert r.status_code == 201, r.text
    assert r.json()["id"] == edge["id"]
    assert r.json()["status"] == "rejected"


# ---------------------------------------------------------------------------
# 3. A sheet rename breaks the join, and confirm-rename repairs it
# ---------------------------------------------------------------------------


async def test_a_renamed_sheet_breaks_the_join_until_the_rename_is_confirmed(
        client, admin_id, tmp_path):
    """FLOW: upload a new version whose sheet was renamed, then repair it.

    Relationships store a logical sheet id and NORMALIZED column names so a
    rename cannot lose them, but the link only follows the sheet once the
    rename is CONFIRMED. This is the whole repair path a UI has to offer: the
    join stops working with a code the screen can recognise, the version diff
    proposes the rename, confirming it makes the SAME relationship id work
    again — under its new name.
    """
    h = auth(admin_id)
    ds = await _sales_dataset(client, admin_id, tmp_path)
    edge = await _declare(client, admin_id, ds)
    assert edge["from_sheet"] == "orders"

    # The join works on v1.
    before = await client.post("/api/v1/joins/preview", headers=h,
                               json={"relationship_id": edge["id"]})
    assert before.status_code == 200, before.text
    assert before.json()["warnings"]["left_rows"] == 3
    columns_before = before.json()["output_columns"]

    # ---- v2 arrives with Orders renamed to Purchases ----------------------
    await _sales_dataset(client, admin_id, tmp_path, name="sales-v2.xlsx",
                         dataset_id=ds, orders_sheet="Purchases")

    # The relationship still lists, still under the identity it was made on.
    listing = await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h)
    assert listing.json()["total"] == 1
    assert listing.json()["items"][0]["from_sheet"] == "orders"

    # ...but the join wizard's default (current version) can no longer find it,
    # and says so in a way a screen can branch on.
    broken = await client.post("/api/v1/joins/preview", headers=h,
                               json={"relationship_id": edge["id"]})
    assert broken.status_code == 404, broken.text
    assert broken.headers["content-type"].startswith(PROBLEM)
    assert broken.json()["code"] == "sheet-not-in-version"
    assert broken.json()["version_number"] == 2

    # ---- The repair path the UI offers: the diff proposes the rename ------
    diff = await client.get(f"/api/v1/datasets/{ds}/versions/1/diff/2", headers=h)
    assert diff.status_code == 200, diff.text
    candidates = [(c["from_sheet"], c["to_sheet"])
                  for c in diff.json()["rename_candidates"]]
    assert ("Orders", "Purchases") in candidates

    confirmed = await client.post(f"/api/v1/datasets/{ds}/versions/2/confirm-rename",
                                  headers=h, json={"from_sheet": "Orders",
                                                   "to_sheet": "Purchases"})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["sheet_key"] == "purchases"

    # ---- The SAME relationship id now reports the NEW sheet name ----------
    listing = await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h)
    assert listing.json()["total"] == 1
    assert listing.json()["items"][0]["id"] == edge["id"]
    assert listing.json()["items"][0]["from_sheet"] == "purchases"
    assert listing.json()["items"][0]["from_column"] == "customer_id"

    detail = await client.get(f"/api/v1/datasets/{ds}/relationships/{edge['id']}",
                              headers=h)
    assert detail.status_code == 200
    assert detail.json()["from_sheet"] == "purchases"
    assert detail.json()["status"] == "confirmed"    # the verdict survived too

    # ---- ...and it drives the join again, on the new version --------------
    after = await client.post("/api/v1/joins/preview", headers=h,
                              json={"relationship_id": edge["id"]})
    assert after.status_code == 200, after.text
    assert after.json()["output_columns"] == columns_before
    assert after.json()["warnings"]["left_rows"] == 3

    run = await client.post("/api/v1/joins/execute", headers=h,
                            json={"relationship_id": edge["id"]})
    assert run.status_code == 200, run.text
    assert run.json()["row_count"] == 3


# ---------------------------------------------------------------------------
# 4. The join wizard
# ---------------------------------------------------------------------------


async def test_the_join_wizard_previews_compares_narrows_executes_and_reads_back(
        client, admin_id, tmp_path):
    """SCREEN: Join builder wizard, step 1 -> 4.

    Step 1 picks a confirmed relationship, step 2 previews and lets the user
    flip inner/left, step 3 ticks output columns, step 4 runs it and shows the
    result. Every number on the warnings panel is asserted, because that panel
    is the entire point of a *guided* join. The last two calls are how the UI
    finds the run it just created.
    """
    h = auth(admin_id)
    ds = await _sales_dataset(client, admin_id, tmp_path, orphan_order=True)
    edge = await _declare(client, admin_id, ds)

    # ---- Step 1: the picker only offers confirmed edges -------------------
    picker = await client.get(f"/api/v1/datasets/{ds}/relationships", headers=h,
                              params={"status": "confirmed"})
    assert picker.status_code == 200, picker.text
    assert picker.json()["total"] == 1
    choice = picker.json()["items"][0]
    assert choice["id"] == edge["id"]
    # The label the wizard renders on the option.
    assert (choice["from_sheet"], choice["from_column"],
            choice["to_sheet"], choice["to_column"]) == (
        "orders", "customer_id", "customers", "customer_id")

    # ---- Step 2a: preview an INNER join -----------------------------------
    inner = await client.post("/api/v1/joins/preview", headers=h,
                              json={"relationship_id": edge["id"], "how": "inner"})
    assert inner.status_code == 200, inner.text
    w = inner.json()["warnings"]
    assert w["left_rows"] == 4 and w["right_rows"] == 3
    assert w["left_duplicate_keys"] == 1     # customer 1 ordered twice
    assert w["right_duplicate_keys"] == 0    # customer_id is a key over there
    assert w["many_to_many"] is False
    assert w["estimated_output_rows"] == 3   # the order for customer 999 drops
    assert w["row_expansion_factor"] == 0.75
    assert w["unmatched_left_pct"] == 25.0
    assert w["unmatched_right_pct"] == 33.33  # customer 3 never ordered
    # The collision banner, and the alias it forces.
    assert w["column_collisions"] == ["region_code"]
    all_columns = inner.json()["output_columns"]
    assert all_columns == ["order_id", "customer_id", "region_code", "total",
                           "customers_region_code", "tier"]
    assert len(inner.json()["preview"]) == 3
    assert set(inner.json()["preview"][0]) == set(all_columns)
    assert inner.json()["relationship"]["id"] == edge["id"]

    # ---- Step 2b: flip to LEFT — the same panel, different numbers --------
    left = await client.post("/api/v1/joins/preview", headers=h,
                             json={"relationship_id": edge["id"], "how": "left"})
    assert left.status_code == 200, left.text
    lw = left.json()["warnings"]
    assert lw["estimated_output_rows"] == 4      # the orphan order is kept
    assert lw["row_expansion_factor"] == 1.0
    assert lw["unmatched_left_pct"] == 25.0      # still reported, not hidden
    assert left.json()["output_columns"] == all_columns
    orphan = [r for r in left.json()["preview"] if r["order_id"] == 13]
    assert len(orphan) == 1 and orphan[0]["tier"] is None

    # ---- Step 3: tick a subset of output columns --------------------------
    picked = ["tier", "order_id", "customers_region_code"]
    narrowed = await client.post("/api/v1/joins/preview", headers=h, json={
        "relationship_id": edge["id"], "how": "left", "select_columns": picked})
    assert narrowed.status_code == 200, narrowed.text
    # NB the response orders columns as the JOIN emits them, not as the user
    # ticked them — a UI that echoes its own order back will disagree.
    assert narrowed.json()["output_columns"] == ["order_id",
                                                 "customers_region_code", "tier"]
    assert set(narrowed.json()["preview"][0]) == set(picked)

    # A typo'd column is caught before anything runs.
    bad = await client.post("/api/v1/joins/preview", headers=h, json={
        "relationship_id": edge["id"], "select_columns": ["region_code", "ghost"]})
    assert bad.status_code == 400, bad.text
    assert bad.json()["code"] == "unknown-column"
    assert "customers_region_code" in bad.json()["available"]

    # ---- Step 4: run it. Preview and execute must agree -------------------
    executed = await client.post("/api/v1/joins/execute", headers=h, json={
        "relationship_id": edge["id"], "how": "left", "select_columns": picked})
    assert executed.status_code == 200, executed.text
    result = executed.json()
    assert result["output_columns"] == narrowed.json()["output_columns"]
    assert result["row_count"] == lw["estimated_output_rows"] == 4
    assert result["warnings"]["column_collisions"] == ["region_code"]
    assert result["relationship"]["id"] == edge["id"]
    run_id, sample_file = result["run_id"], result["sample_file"]

    # ---- The result grid reads the artifact the run produced --------------
    data = await client.get(f"/api/v1/samples/{sample_file}/data", headers=h)
    assert data.status_code == 200, data.text
    assert data.json()["total_count"] == 4
    assert [c["name"] for c in data.json()["columns"]] == result["output_columns"]
    assert len(data.json()["data"]) == 4
    assert {row["order_id"] for row in data.json()["data"]} == {10, 11, 12, 13}

    # ---- "View in library": the only way to reach the run from here -------
    # /joins/execute returns run_id but no definition_id, and there is no
    # GET /runs/{run_id}, so the UI has to scan the definition list first.
    defs = await client.get(f"/api/v1/datasets/{ds}/analytics", headers=h)
    assert defs.status_code == 200, defs.text
    assert defs.json()["total"] == 1
    definition = defs.json()["items"][0]
    assert definition["kind"] == "join"
    assert definition["name"] == f"join:{edge['id'][:8]}:left"

    runs = await client.get(
        f"/api/v1/datasets/{ds}/analytics/{definition['id']}/runs", headers=h)
    assert runs.status_code == 200, runs.text
    assert runs.json()["total"] == 1
    assert runs.json()["items"][0]["id"] == run_id
    assert runs.json()["items"][0]["status"] == "completed"


# ---------------------------------------------------------------------------
# 5. Version pinning
# ---------------------------------------------------------------------------


async def test_a_join_pinned_to_an_older_version_survives_the_data_moving_on(
        client, admin_id, tmp_path):
    """FLOW: reproduce a historical join after the dataset has changed shape.

    ``left_version``/``right_version`` are the only way a UI reproduces a join
    someone ran last quarter. The two ways the current version can betray a
    relationship — the key column dropped, the whole sheet dropped — return
    DIFFERENT codes, and a screen must render different repairs for them
    ("re-map the column" vs "this sheet is gone").
    """
    h = auth(admin_id)
    ds = await _sales_dataset(client, admin_id, tmp_path)
    edge = await _declare(client, admin_id, ds)

    # ---- v1: the default and the explicit pin agree -----------------------
    default_preview = await client.post("/api/v1/joins/preview", headers=h,
                                        json={"relationship_id": edge["id"]})
    assert default_preview.status_code == 200, default_preview.text
    pinned = await client.post("/api/v1/joins/preview", headers=h, json={
        "relationship_id": edge["id"], "left_version": 1, "right_version": 1})
    assert pinned.status_code == 200, pinned.text
    assert pinned.json()["output_columns"] == default_preview.json()["output_columns"]
    assert pinned.json()["warnings"] == default_preview.json()["warnings"]

    # ---- v2 drops the key column from Orders ------------------------------
    await _sales_dataset(client, admin_id, tmp_path, name="sales-v2.xlsx",
                         dataset_id=ds, orders_key=False)

    broken = await client.post("/api/v1/joins/preview", headers=h,
                               json={"relationship_id": edge["id"]})
    assert broken.status_code == 400, broken.text
    assert broken.headers["content-type"].startswith(PROBLEM)
    assert broken.json()["code"] == "relationship-endpoint-mismatch"
    assert broken.json()["column"] == "customer_id"
    assert broken.json()["sheet"] == "Orders"

    # The pin still works — that is the whole point of the field.
    still_good = await client.post("/api/v1/joins/preview", headers=h, json={
        "relationship_id": edge["id"], "left_version": 1})
    assert still_good.status_code == 200, still_good.text
    assert still_good.json()["warnings"]["left_rows"] == 3

    executed = await client.post("/api/v1/joins/execute", headers=h, json={
        "relationship_id": edge["id"], "left_version": 1, "right_version": 1})
    assert executed.status_code == 200, executed.text
    assert executed.json()["row_count"] == 3
    run_id = executed.json()["run_id"]

    defs = await client.get(f"/api/v1/datasets/{ds}/analytics", headers=h)
    definition = next(d for d in defs.json()["items"] if d["kind"] == "join")
    runs = await client.get(
        f"/api/v1/datasets/{ds}/analytics/{definition['id']}/runs", headers=h)
    assert [r["id"] for r in runs.json()["items"]] == [run_id]

    # ---- v3 drops the whole Orders sheet: a DIFFERENT failure -------------
    await _sales_dataset(client, admin_id, tmp_path, name="sales-v3.xlsx",
                         dataset_id=ds, include_orders=False)

    gone = await client.post("/api/v1/joins/preview", headers=h,
                             json={"relationship_id": edge["id"]})
    assert gone.status_code == 404, gone.text
    assert gone.json()["code"] == "sheet-not-in-version"
    assert gone.json()["version_number"] == 3

    # A version that never existed is a plain 404 with no relationship code.
    nowhere = await client.post("/api/v1/joins/preview", headers=h, json={
        "relationship_id": edge["id"], "left_version": 99})
    assert nowhere.status_code == 404, nowhere.text

    # The relationship itself is untouched by any of this, so the UI can still
    # offer the pinned re-run.
    detail = await client.get(f"/api/v1/datasets/{ds}/relationships/{edge['id']}",
                              headers=h)
    assert detail.json()["status"] == "confirmed"
    replay = await client.post("/api/v1/joins/preview", headers=h, json={
        "relationship_id": edge["id"], "left_version": 1, "right_version": 1})
    assert replay.status_code == 200, replay.text
    assert replay.json()["warnings"]["estimated_output_rows"] == 3


# ---------------------------------------------------------------------------
# 6. Publishing the join output back onto the left dataset
# ---------------------------------------------------------------------------


async def test_a_join_output_is_published_as_a_new_version_and_lineage_names_both_parents(
        client, admin_id, tmp_path):
    """SCREEN: join result > "Publish" > "as a new version of ...".

    ``mode=new_version`` writes the joined table over the left dataset's
    workbook shape, so the version list, the version preview and the lineage
    panel all change at once. A UI showing "derived from" needs both parents,
    and a UI showing the dataset needs to know its sheets were replaced.
    """
    h = auth(admin_id)
    left = await _sales_dataset(client, admin_id, tmp_path, name="left.xlsx")
    right = await _sales_dataset(client, admin_id, tmp_path, name="right.xlsx")
    edge = await _declare(client, admin_id, left, to_dataset_id=right)
    assert edge["to_dataset_id"] == right

    executed = await client.post("/api/v1/joins/execute", headers=h,
                                 json={"relationship_id": edge["id"]})
    assert executed.status_code == 200, executed.text
    run_id = executed.json()["run_id"]
    joined_columns = executed.json()["output_columns"]
    row_count = executed.json()["row_count"]

    # ---- Publish onto the LEFT dataset (the default target) ---------------
    published = await client.post(f"/api/v1/joins/{run_id}/publish", headers=h,
                                  json={"mode": "new_version"})
    assert published.status_code == 200, published.text
    body = published.json()
    assert body["dataset_id"] == left
    assert body["mode"] == "new_version"
    assert body["version_number"] == 2
    assert body["version_id"] and body["dataset_name"]

    # ---- The version list the screen refreshes to -------------------------
    versions = await client.get(f"/api/v1/datasets/{left}/versions", headers=h)
    assert versions.status_code == 200, versions.text
    assert [v["version_number"] for v in versions.json()["items"]] == [2, 1]
    assert versions.json()["items"][0]["id"] == body["version_id"]
    assert versions.json()["items"][0]["row_count"] == row_count

    # ---- The joined table REPLACED the three-sheet workbook --------------
    meta = await client.get(f"/api/v1/datasets/{left}", headers=h)
    assert meta.status_code == 200, meta.text
    assert len(meta.json()["sheets"]) == 1, "a published join is a single table"

    preview = await client.get(f"/api/v1/datasets/{left}/versions/2/preview",
                               headers=h)
    assert preview.status_code == 200, preview.text
    assert preview.json()["total"] == row_count
    assert set(preview.json()["items"][0]) == set(joined_columns)
    assert all("tier" in row for row in preview.json()["items"])

    # ---- Lineage: two parents, each against a version of its OWN dataset --
    lineage = await client.get(f"/api/v1/datasets/{left}/lineage", headers=h)
    assert lineage.status_code == 200, lineage.text
    parents = [p for p in lineage.json()["parents"]
               if p["dataset_version_id"] == body["version_id"]]
    assert len(parents) == 2
    assert {p["relation"] for p in parents} == {"joined_from"}
    assert {p["parent_dataset_id"] for p in parents} == {left, right}
    for parent in parents:
        owner = parent["parent_dataset_id"]
        owned = {v["id"] for v in (await client.get(
            f"/api/v1/datasets/{owner}/versions", headers=h)).json()["items"]}
        assert parent["parent_version_id"] in owned, parent
        assert parent["parent_version_number"] == 1

    # ---- Publishing the same run twice appends, it does not overwrite -----
    again = await client.post(f"/api/v1/joins/{run_id}/publish", headers=h,
                              json={"mode": "new_version"})
    assert again.status_code == 200, again.text
    assert again.json()["version_number"] == 3
    versions = await client.get(f"/api/v1/datasets/{left}/versions", headers=h)
    assert [v["version_number"] for v in versions.json()["items"]] == [3, 2, 1]

    # ---- And the other mode still makes a standalone dataset -------------
    fresh = await client.post(f"/api/v1/joins/{run_id}/publish", headers=h,
                              json={"mode": "new_dataset",
                                    "name": f"joined-{run_id[:8]}"})
    assert fresh.status_code == 200, fresh.text
    assert fresh.json()["dataset_id"] != left
    assert fresh.json()["version_number"] == 1
    assert fresh.json()["dataset_name"] == f"joined-{run_id[:8]}"


# ---------------------------------------------------------------------------
# 7. Deleting the edge behind a saved join
# ---------------------------------------------------------------------------


async def test_deleting_the_relationship_behind_a_saved_join_is_refused_until_cleared(
        client, admin_id, tmp_path):
    """FLOW: Relationships > Delete, seen from the analytics library afterwards.

    The delete is refused with a 409 that NAMES the saved join in the way, and
    only goes through once that library row is gone. Without the guard the
    library was left holding a definition whose `params.relationship_id` points
    at nothing: it still lists, still counts towards `total`, still shows its
    run history — and Re-run/Preview/Publish on that row all answer 404, a
    dead end the screen has no way to predict.
    """
    h = auth(admin_id)
    ds = await _sales_dataset(client, admin_id, tmp_path)
    edge = await _declare(client, admin_id, ds)

    executed = await client.post("/api/v1/joins/execute", headers=h,
                                 json={"relationship_id": edge["id"]})
    assert executed.status_code == 200, executed.text
    run_id, sample_file = executed.json()["run_id"], executed.json()["sample_file"]

    defs = await client.get(f"/api/v1/datasets/{ds}/analytics", headers=h)
    definition = next(d for d in defs.json()["items"] if d["kind"] == "join")
    assert definition["params"]["relationship_id"] == edge["id"]

    # ---- Delete, from the relationship detail screen — refused ------------
    r = await client.delete(f"/api/v1/datasets/{ds}/relationships/{edge['id']}",
                            headers=h)
    assert r.status_code == 409, r.text
    assert r.headers["content-type"].startswith(PROBLEM)
    assert r.json()["code"] == "relationship-has-dependents"
    # The dialog can list exactly what is in the way, by id AND by name.
    blocking = r.json()["attached"]["join_definitions"]
    assert [d["id"] for d in blocking] == [definition["id"]]
    assert blocking[0]["name"] == definition["name"]

    # Nothing moved: the edge and everything hanging off it still work.
    assert (await client.get(f"/api/v1/datasets/{ds}/relationships/{edge['id']}",
                             headers=h)).status_code == 200
    assert (await client.get(f"/api/v1/datasets/{ds}/relationships",
                             headers=h)).json()["total"] == 1
    assert (await client.post("/api/v1/joins/preview", headers=h,
                              json={"relationship_id": edge["id"]})
            ).status_code == 200

    # ---- The library still shows the join, with its history intact --------
    defs = await client.get(f"/api/v1/datasets/{ds}/analytics", headers=h)
    assert defs.json()["total"] == 1
    assert defs.json()["items"][0]["id"] == definition["id"]

    runs = await client.get(
        f"/api/v1/datasets/{ds}/analytics/{definition['id']}/runs", headers=h)
    assert runs.json()["total"] == 1
    assert runs.json()["items"][0]["id"] == run_id
    assert runs.json()["items"][0]["status"] == "completed"

    # The data it produced is still there and still readable.
    data = await client.get(f"/api/v1/samples/{sample_file}/data", headers=h)
    assert data.status_code == 200, data.text
    assert data.json()["total_count"] == 3

    # ---- Clear the dependent, and the delete goes through -----------------
    assert (await client.delete(
        f"/api/v1/datasets/{ds}/analytics/{definition['id']}",
        headers=h)).status_code == 204
    r = await client.delete(f"/api/v1/datasets/{ds}/relationships/{edge['id']}",
                            headers=h)
    assert r.status_code == 204, r.text
    assert not r.content

    # Deleting twice is a 404, so a double-click surfaces cleanly.
    r = await client.delete(f"/api/v1/datasets/{ds}/relationships/{edge['id']}",
                            headers=h)
    assert r.status_code == 404, r.text

    assert (await client.get(f"/api/v1/datasets/{ds}/relationships/{edge['id']}",
                             headers=h)).status_code == 404
    assert (await client.get(f"/api/v1/datasets/{ds}/relationships",
                             headers=h)).json()["total"] == 0

    # ---- ...and now every join route for that edge is a clean 404 ---------
    rerun = await client.post("/api/v1/joins/execute", headers=h,
                              json={"relationship_id": edge["id"]})
    assert rerun.status_code == 404, rerun.text
    preview = await client.post("/api/v1/joins/preview", headers=h,
                                json={"relationship_id": edge["id"]})
    assert preview.status_code == 404, preview.text
    publish = await client.post(f"/api/v1/joins/{run_id}/publish", headers=h,
                                json={"mode": "new_dataset", "name": "orphan"})
    assert publish.status_code == 404, publish.text


# ---------------------------------------------------------------------------
# 8. The permission matrix
# ---------------------------------------------------------------------------


async def test_the_relationships_screen_permission_matrix_for_viewer_editor_and_outsider(
        client, admin_id, tmp_path):
    """SCREEN: Relationships + join builder, rendered for three roles.

    A read-only screen must know which controls to hide, and a cross-team user
    must not learn the dataset exists at all. The split the UI depends on:
    in-team-but-wrong-role is 403 (hide the button), out-of-team is 404 (hide
    the whole route).
    """
    h = auth(admin_id)
    ds = await _sales_dataset(client, admin_id, tmp_path)
    confirmed_edge = await _declare(client, admin_id, ds)
    suggested_edge = await _declare(client, admin_id, ds, from_sheet="Orders",
                                    from_column="region_code",
                                    to_sheet="Regions", to_column="region_code",
                                    confirmed=False)

    viewer, _ = await create_team_user(client, admin_id, "viewer",
                                       team_id=DEFAULT_TEAM_ID)
    editor, _ = await create_team_user(client, admin_id, "editor",
                                       team_id=DEFAULT_TEAM_ID)
    outsider, _ = await create_team_user(client, admin_id, "admin")
    vh, eh, oh = auth(viewer), auth(editor), auth(outsider)
    base = f"/api/v1/datasets/{ds}/relationships"

    # ---- Viewer: the screen renders, read-only ---------------------------
    listing = await client.get(base, headers=vh)
    assert listing.status_code == 200, listing.text
    assert listing.json()["total"] == 2
    detail = await client.get(f"{base}/{confirmed_edge['id']}", headers=vh)
    assert detail.status_code == 200, detail.text
    assert detail.json()["status"] == "confirmed"

    # A viewer may explore what a join WOULD do — nothing is persisted.
    explore = await client.post("/api/v1/joins/preview", headers=vh,
                                json={"relationship_id": confirmed_edge["id"]})
    assert explore.status_code == 200, explore.text
    assert explore.json()["warnings"]["estimated_output_rows"] == 3

    # ...and every control that changes state is a 403 the UI can hide on.
    forbidden = [
        ("POST", f"{base}/seed", None),
        ("POST", f"{base}/suggest", None),
        ("POST", base, {"from_sheet": "Orders", "from_column": "customer_id",
                        "to_sheet": "Customers", "to_column": "customer_id"}),
        ("POST", f"{base}/{suggested_edge['id']}/confirm", None),
        ("POST", f"{base}/{suggested_edge['id']}/reject", None),
        ("DELETE", f"{base}/{suggested_edge['id']}", None),
        # Executing is a write: it leaves behind a definition, a job, a run and
        # a parquet artifact. Only the PREVIEW above is open to a viewer.
        ("POST", "/api/v1/joins/execute", {"relationship_id": confirmed_edge["id"]}),
    ]
    for method, path, payload in forbidden:
        r = await client.request(method, path, headers=vh, json=payload)
        assert r.status_code == 403, (method, path, r.status_code, r.text)

    # Nothing the viewer tried changed anything.
    assert (await client.get(base, headers=h)).json()["total"] == 2
    assert (await client.get(f"{base}/{suggested_edge['id']}",
                             headers=h)).json()["status"] == "suggested"
    # Including the refused execute: no library row was materialized.
    assert (await client.get(f"/api/v1/datasets/{ds}/analytics",
                             headers=h)).json()["total"] == 0

    # ---- Editor: the same screen, controls enabled ------------------------
    reviewed = await client.post(f"{base}/{suggested_edge['id']}/confirm", headers=eh)
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["reviewed_by"] == editor

    declared = await client.post(base, headers=eh, json={
        "from_sheet": "Customers", "from_column": "region_code",
        "to_sheet": "Regions", "to_column": "region_code"})
    assert declared.status_code == 201, declared.text

    ran = await client.post("/api/v1/joins/execute", headers=eh,
                            json={"relationship_id": confirmed_edge["id"]})
    assert ran.status_code == 200, ran.text
    run_id = ran.json()["run_id"]

    # An editor may publish onto a dataset in their own team...
    pub = await client.post(f"/api/v1/joins/{run_id}/publish", headers=eh,
                            json={"mode": "new_dataset", "name": f"e-{run_id[:8]}"})
    assert pub.status_code == 200, pub.text

    # ...but a viewer may not, even on a run they can otherwise read.
    denied = await client.post(f"/api/v1/joins/{run_id}/publish", headers=vh,
                               json={"mode": "new_dataset", "name": "nope"})
    assert denied.status_code == 403, denied.text

    assert (await client.delete(f"{base}/{declared.json()['id']}",
                                headers=eh)).status_code == 204

    # ---- Outsider: the routes do not exist for them -----------------------
    hidden = [
        ("GET", base, None),
        ("GET", f"{base}/{confirmed_edge['id']}", None),
        ("POST", base, {"from_sheet": "Orders", "from_column": "customer_id",
                        "to_sheet": "Customers", "to_column": "customer_id"}),
        ("POST", f"{base}/seed", None),
        ("POST", f"{base}/suggest", None),
        ("POST", f"{base}/{confirmed_edge['id']}/confirm", None),
        ("POST", f"{base}/{confirmed_edge['id']}/reject", None),
        ("DELETE", f"{base}/{confirmed_edge['id']}", None),
        ("POST", "/api/v1/joins/preview", {"relationship_id": confirmed_edge["id"]}),
        ("POST", "/api/v1/joins/execute", {"relationship_id": confirmed_edge["id"]}),
        ("POST", f"/api/v1/joins/{run_id}/publish", {"mode": "new_dataset",
                                                     "name": "sneak"}),
    ]
    for method, path, payload in hidden:
        r = await client.request(method, path, headers=oh, json=payload)
        assert r.status_code == 404, (method, path, r.status_code, r.text)
        assert ds not in r.text or "not found" in r.text.lower()
