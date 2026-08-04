"""Phase 1 integration tests — first-class sheets, diffs, tag promotion.

Covers: per-sheet schema capture (normalized columns, fingerprints, hidden
sheets), the sheet-selection-required contract, workbook + sheet diffs, and
explicit tag promote/rollback/history with reason + actor.
"""

from __future__ import annotations

import pytest_asyncio
from openpyxl import Workbook

from conftest import DEFAULT_TEAM_ID, auth

PROBLEM = "application/problem+json"


def _workbook_v1(path):
    """Three sheets: dup/blank headers, a second sheet, and a hidden sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Amount", "Amount", None, "Region"])
    for row in ([100, 1, "a", "EU"], [200, 2, "b", "US"], [300, 3, "c", "APAC"]):
        ws.append(row)
    exp = wb.create_sheet("Expenses")
    exp.append(["Item", "Cost"])
    exp.append(["rent", 50])
    exp.append(["power", 20])
    sec = wb.create_sheet("Secrets")
    sec.append(["K", "V"])
    sec.append(["k1", "v1"])
    sec.sheet_state = "hidden"
    wb.save(path)


def _workbook_v2(path):
    """Revenue gains a column + a row; Expenses renamed to Spending; Secrets same."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Amount", "Amount", None, "Region", "Quarter"])
    for row in ([100, 1, "a", "EU", "Q1"], [200, 2, "b", "US", "Q1"],
                [300, 3, "c", "APAC", "Q2"], [400, 4, "d", "EU", "Q2"]):
        ws.append(row)
    sp = wb.create_sheet("Spending")
    sp.append(["Item", "Cost"])
    sp.append(["rent", 55])
    sp.append(["power", 25])
    sec = wb.create_sheet("Secrets")
    sec.append(["K", "V"])
    sec.append(["k1", "v1"])
    sec.sheet_state = "hidden"
    wb.save(path)


async def _upload_xlsx(client, admin_id, path, name, dataset_id=None):
    data = {"dataset_id": dataset_id} if dataset_id else {}
    with open(path, "rb") as f:
        r = await client.post(
            "/api/v1/upload",
            headers={**auth(admin_id), "X-Team-Id": DEFAULT_TEAM_ID},
            files={"file": (name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data=data,
        )
    assert r.status_code == 200, r.text
    return r.json()["dataset_id"]


@pytest_asyncio.fixture
async def excel_dataset(client, admin_id, tmp_path):
    """A dataset with two Excel versions (v1 and v2 workbooks above)."""
    v1 = tmp_path / "v1.xlsx"
    v2 = tmp_path / "v2.xlsx"
    _workbook_v1(v1)
    _workbook_v2(v2)
    ds = await _upload_xlsx(client, admin_id, v1, "book.xlsx")
    await _upload_xlsx(client, admin_id, v2, "book.xlsx", dataset_id=ds)
    return ds


# --- Sheets are first-class ---------------------------------------------------

async def test_sheet_rows_capture_schema_and_visibility(client, admin_id, excel_dataset):
    r = await client.get(f"/api/v1/datasets/{excel_dataset}/sheets", headers=auth(admin_id))
    assert r.status_code == 200, r.text
    sheets = {s["name"]: s for s in r.json()["items"]}
    assert set(sheets) == {"Revenue", "Spending", "Secrets"}  # current version = v2

    assert sheets["Secrets"]["visibility"] == "hidden"
    assert sheets["Revenue"]["visibility"] == "visible"
    for s in sheets.values():
        assert s["schema_fingerprint"], s["name"]
        assert s["sheet_key"], s["name"]

    # Duplicate + blank headers: original names kept, normalized names unique.
    rev_cols = {c["normalized_name"]: c for c in sheets["Revenue"]["columns"]}
    assert {"amount", "amount_2", "column_2", "region", "quarter"} == set(rev_cols)
    assert rev_cols["amount"]["original_name"] == "Amount"
    assert rev_cols["amount_2"]["original_name"] == "Amount"
    assert rev_cols["column_2"]["original_name"] == ""


async def test_csv_dataset_gets_synthetic_data_sheet(client, admin_id, admin_dataset):
    r = await client.get(f"/api/v1/datasets/{admin_dataset}/sheets", headers=auth(admin_id))
    assert r.status_code == 200
    items = r.json()["items"]
    assert len(items) == 1 and items[0]["name"] == "data"
    assert items[0]["schema_fingerprint"]
    assert all(c["normalized_name"] and c["dtype"] for c in items[0]["columns"])


async def test_multi_sheet_requires_explicit_sheet(client, admin_id, excel_dataset):
    h = auth(admin_id)
    # No sheet named → machine-readable problem listing the choices.
    r = await client.post("/api/v1/sample", headers=h, json={
        "dataset_id": excel_dataset, "target_total_volume": 2,
        "sampling_steps": [{"method": "random", "sample_size": 2}],
    })
    assert r.status_code == 400, r.text
    assert r.headers["content-type"].startswith(PROBLEM)
    body = r.json()
    assert body["code"] == "sheet-selection-required"
    assert set(body["sheets"]) == {"Revenue", "Spending", "Secrets"}

    # Same for downloads.
    r = await client.get(f"/api/v1/datasets/{excel_dataset}/download",
                         params={"format": "csv"}, headers=h)
    assert r.status_code == 400 and r.json()["code"] == "sheet-selection-required"

    # Naming a sheet works.
    r = await client.post("/api/v1/sample", headers=h, json={
        "dataset_id": excel_dataset, "sheet": "Spending", "target_total_volume": 2,
        "sampling_steps": [{"method": "random", "sample_size": 2}],
    })
    assert r.status_code == 200, r.text

    r = await client.get(f"/api/v1/datasets/{excel_dataset}/download",
                         params={"format": "csv", "sheet": "Spending"}, headers=h)
    assert r.status_code == 200
    assert b"Item" in r.content

    # Version-pinned downloads follow the same contract.
    r = await client.get(f"/api/v1/datasets/{excel_dataset}/versions/1/download",
                         params={"format": "parquet"}, headers=h)
    assert r.status_code == 400 and r.json()["code"] == "sheet-selection-required"
    r = await client.get(f"/api/v1/datasets/{excel_dataset}/versions/1/download",
                         params={"format": "parquet", "sheet": "Expenses"}, headers=h)
    assert r.status_code == 200 and r.content[:4] == b"PAR1"


# --- Diffs -------------------------------------------------------------------

async def test_workbook_diff(client, admin_id, excel_dataset):
    r = await client.get(f"/api/v1/datasets/{excel_dataset}/versions/1/diff/2",
                         headers=auth(admin_id))
    assert r.status_code == 200, r.text
    body = r.json()
    assert [s["name"] for s in body["added"]] == ["Spending"]
    assert [s["name"] for s in body["removed"]] == ["Expenses"]
    assert body["unchanged"] == ["Secrets"]
    modified = {m["sheet_key"]: m for m in body["modified"]}
    assert set(modified) == {"revenue"}
    assert modified["revenue"]["schema_changed"] is True
    assert modified["revenue"]["row_count_delta"] == 1
    # Rename is only ever a suggestion.
    cands = body["rename_candidates"]
    assert len(cands) == 1
    assert (cands[0]["from_sheet"], cands[0]["to_sheet"]) == ("Expenses", "Spending")
    assert cands[0]["confidence"] == "high"


async def test_sheet_diff(client, admin_id, excel_dataset):
    r = await client.get(
        f"/api/v1/datasets/{excel_dataset}/versions/1/sheets/Revenue/diff/2",
        headers=auth(admin_id))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["identical"] is False
    assert [c["normalized_name"] for c in body["added_columns"]] == ["quarter"]
    assert body["removed_columns"] == []
    assert body["row_count_delta"] == 1

    # Unchanged sheet diffs as identical.
    r = await client.get(
        f"/api/v1/datasets/{excel_dataset}/versions/1/sheets/Secrets/diff/2",
        headers=auth(admin_id))
    assert r.status_code == 200 and r.json()["identical"] is True

    # A sheet missing from the target version 404s.
    r = await client.get(
        f"/api/v1/datasets/{excel_dataset}/versions/1/sheets/Expenses/diff/2",
        headers=auth(admin_id))
    assert r.status_code == 404


async def test_sheet_diff_on_csv_style_versions(client, admin_id):
    """Inline JSON versions diff via the synthetic 'data' sheet."""
    h = auth(admin_id)
    r = await client.post("/api/v1/upload", headers={**h, "X-Team-Id": DEFAULT_TEAM_ID},
                          data={"data": '[{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]'})
    assert r.status_code == 200, r.text
    ds = r.json()["dataset_id"]
    r = await client.post("/api/v1/upload", headers=h,
                          data={"data": '[{"a": 1, "c": 5}]', "dataset_id": ds})
    assert r.status_code == 200, r.text

    r = await client.get(f"/api/v1/datasets/{ds}/versions/1/sheets/data/diff/2", headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert [c["normalized_name"] for c in body["added_columns"]] == ["c"]
    assert [c["normalized_name"] for c in body["removed_columns"]] == ["b"]
    assert body["row_count_delta"] == -1


# --- Tag promotion lifecycle --------------------------------------------------

async def test_tag_promote_rollback_history(client, admin_id, excel_dataset):
    h = auth(admin_id)
    ds = excel_dataset
    # Promote to v1 with a reason, then to v2.
    r = await client.post(f"/api/v1/datasets/{ds}/tags/production/promote", headers=h,
                          json={"version_number": 1, "reason": "initial go-live"})
    assert r.status_code == 200, r.text
    assert r.json()["to_version_number"] == 1 and r.json()["from_version_number"] is None
    r = await client.post(f"/api/v1/datasets/{ds}/tags/production/promote", headers=h,
                          json={"version_number": 2, "reason": "monthly refresh"})
    assert r.status_code == 200
    assert r.json()["from_version_number"] == 1 and r.json()["to_version_number"] == 2

    # Rollback returns to v1 and the tag resolves there.
    r = await client.post(f"/api/v1/datasets/{ds}/tags/production/rollback", headers=h,
                          json={"reason": "bad refresh"})
    assert r.status_code == 200, r.text
    assert r.json()["from_version_number"] == 2 and r.json()["to_version_number"] == 1
    r = await client.get(f"/api/v1/datasets/{ds}/tags/production", headers=h)
    assert r.status_code == 200 and r.json()["version_number"] == 1

    # History: newest first, with actions, reasons, and the actor.
    r = await client.get(f"/api/v1/datasets/{ds}/tags/production/history", headers=h)
    assert r.status_code == 200
    entries = r.json()["items"]
    assert [e["action"] for e in entries] == ["rollback", "promote", "promote"]
    assert entries[0]["reason"] == "bad refresh"
    assert entries[0]["from_version_number"] == 2 and entries[0]["to_version_number"] == 1
    assert entries[0]["actor_email"] == "system@localhost"


async def test_tag_promote_validation_and_rollback_guard(client, admin_id, excel_dataset):
    h = auth(admin_id)
    ds = excel_dataset
    # Unknown target version → 404; missing target → 400.
    r = await client.post(f"/api/v1/datasets/{ds}/tags/staging/promote", headers=h,
                          json={"version_number": 99})
    assert r.status_code == 404
    r = await client.post(f"/api/v1/datasets/{ds}/tags/staging/promote", headers=h, json={})
    assert r.status_code == 400

    # A tag that only ever pointed at one version can't roll back.
    r = await client.post(f"/api/v1/datasets/{ds}/tags/staging/promote", headers=h,
                          json={"version_number": 2})
    assert r.status_code == 200
    r = await client.post(f"/api/v1/datasets/{ds}/tags/staging/rollback", headers=h, json={})
    assert r.status_code == 409

    # Raw PUT and DELETE also land in history.
    r = await client.put(f"/api/v1/datasets/{ds}/tags", headers=h,
                         json={"tag_name": "staging", "version_number": 1})
    assert r.status_code == 200
    r = await client.delete(f"/api/v1/datasets/{ds}/tags/staging", headers=h)
    assert r.status_code == 200
    r = await client.get(f"/api/v1/datasets/{ds}/tags/staging/history", headers=h)
    assert r.status_code == 200
    actions = [e["action"] for e in r.json()["items"]]
    assert actions == ["delete", "set", "promote"]
