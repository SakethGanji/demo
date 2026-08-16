"""Wave 2 §12 — exporting stored result files to CSV/XLSX/Parquet.

Exports are owned `export` artifacts (ownership inherited from the source
file), so the standing /samples authorization contracts hold for them too.
"""

from __future__ import annotations

import json

from openpyxl import load_workbook

from conftest import auth, create_team_user, upload_inline

ROWS = [
    {"region": "EU", "amount": 100.0},
    {"region": "US", "amount": 50.0},
    {"region": "EU", "amount": 200.0},
]


async def _team_pivot_file(client, admin_id):
    """An editor in a fresh team makes a pivot; returns (editor, team, filename)."""
    editor, team = await create_team_user(client, admin_id, "editor")
    ds = (await upload_inline(client, editor, json.dumps(ROWS),
                              team_id=team))["dataset_id"]
    r = await client.post("/api/v1/pivot", headers=auth(editor), json={
        "dataset_id": ds, "rows": ["region"],
        "values": [{"column": "amount", "function": "sum", "alias": "amt"}]})
    assert r.status_code == 200, r.text
    return editor, team, r.json()["result_file"]


async def test_export_csv_and_xlsx_roundtrip(client, admin_id, tmp_path):
    editor, _team, fname = await _team_pivot_file(client, admin_id)
    h = auth(editor)

    r = await client.post(f"/api/v1/samples/{fname}/export",
                          params={"format": "csv"}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["format"] == "csv" and body["source_file"] == fname
    assert body["export_file"].endswith(".csv")

    r = await client.get(f"/api/v1/samples/{body['export_file']}", headers=h)
    assert r.status_code == 200
    text = r.content.decode()
    assert text.splitlines()[0] == "region,amt"
    assert "EU,300.0" in text

    r = await client.post(f"/api/v1/samples/{fname}/export",
                          params={"format": "xlsx"}, headers=h)
    assert r.status_code == 200, r.text
    xlsx_name = r.json()["export_file"]
    r = await client.get(f"/api/v1/samples/{xlsx_name}", headers=h)
    assert r.status_code == 200
    local = tmp_path / "export.xlsx"
    local.write_bytes(r.content)
    ws = load_workbook(local).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0] == ["region", "amt"]
    assert ["EU", 300.0] in rows


async def test_export_authorization_inherited(client, admin_id):
    editor, team, fname = await _team_pivot_file(client, admin_id)
    export = (await client.post(f"/api/v1/samples/{fname}/export",
                                params={"format": "parquet"},
                                headers=auth(editor))).json()["export_file"]

    # Same-team viewer reads the export; outsiders get existence-hiding 404s.
    viewer, _ = await create_team_user(client, admin_id, "viewer", team_id=team)
    r = await client.get(f"/api/v1/samples/{export}", headers=auth(viewer))
    assert r.status_code == 200

    outsider, _ = await create_team_user(client, admin_id, "editor")
    for url in (f"/api/v1/samples/{export}",
                f"/api/v1/samples/{fname}/export?format=csv"):
        method = client.get if "export?" not in url else client.post
        r = await method(url, headers=auth(outsider))
        assert r.status_code == 404, url


async def test_export_error_contracts(client, admin_id):
    editor, _team, fname = await _team_pivot_file(client, admin_id)
    h = auth(editor)

    r = await client.post(f"/api/v1/samples/{fname}/export",
                          params={"format": "yaml"}, headers=h)
    assert r.status_code == 400 and "Unsupported export format" in r.json()["detail"]

    r = await client.post("/api/v1/samples/nope.parquet/export",
                          params={"format": "csv"}, headers=auth(admin_id))
    assert r.status_code == 404
