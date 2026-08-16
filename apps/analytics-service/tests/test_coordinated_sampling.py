"""Coordinated cross-sheet sampling — POST /api/v1/sample/coordinated.

Sample a driver sheet, then semi-join related sheets down to the rows the
sample references (driver by default, chained via ``parent_sheet``). Keys are
declared explicitly, resolved from a confirmed relationship (§22), or defaulted
from a foreign_key quality rule (§5) — in that precedence. A related sheet may
additionally be sub-sampled (§24), except when another link depends on it.
"""

from __future__ import annotations

from openpyxl import Workbook

from conftest import auth, create_team_user, make_orders_workbook, upload_file

PROBLEM = "application/problem+json"


def _chain_workbook(path):
    """Regions ← Customers ← Orders: a two-hop FK chain for parent_sheet tests."""
    wb = Workbook()
    reg = wb.active
    reg.title = "Regions"
    reg.append(["region_id", "region_name"])
    for r in ([1, "north"], [2, "south"], [3, "east"]):
        reg.append(r)
    cust = wb.create_sheet("Customers")
    cust.append(["customer_id", "region_id", "tier"])
    for r in ([1, 1, "gold"], [2, 1, "silver"], [3, 2, "gold"],
              [4, 3, "bronze"], [5, 3, "gold"]):
        cust.append(r)
    orders = wb.create_sheet("Orders")
    orders.append(["order_id", "customer_id", "total"])
    for r in ([10, 1, 9.5], [11, 1, 12.0], [12, 2, 3.0],
              [13, 4, 88.0], [14, 5, 1.0]):
        orders.append(r)
    wb.save(path)


async def _upload_orders(client, admin_id, tmp_path):
    p = tmp_path / "orders.xlsx"
    make_orders_workbook(p, clean=True)
    return (await upload_file(client, admin_id, p))["dataset_id"]


def _coord_body(ds, **over):
    body = {
        "dataset_id": ds, "driver_sheet": "Orders", "target_total_volume": 2,
        "sampling_steps": [{"method": "random", "sample_size": 2}],
        "seed": 42,
        "related": [{"sheet": "Customers", "left_on": "customer_id",
                     "right_on": "customer_id"}],
    }
    body.update(over)
    return body


async def test_referential_consistency_and_reduction(client, admin_id, tmp_path):
    """The related sheet is filtered to exactly the keys the driver sampled —
    and the filter genuinely *reduces* it (3 customers, only 2 referenced)."""
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)

    r = await client.post("/api/v1/sample/coordinated", headers=h,
                          json=_coord_body(ds))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["success"] and body["driver_sheet"] == "Orders"

    driver_keys = {row["customer_id"] for row in body["driver"]["data"]}
    assert body["driver"]["sampled_count"] == 2

    rel = body["related"][0]
    assert rel["sheet"] == "Customers" and rel["parent_sheet"] == "Orders"
    assert rel["original_count"] == 3
    rel_keys = {row["customer_id"] for row in rel["data"]}
    assert rel_keys == driver_keys          # exact semi-join, no strays
    assert rel["sampled_count"] < rel["original_count"]  # actually reduced
    assert rel["columns"] and rel["preview"]


async def test_chain_via_parent_sheet_resolved_out_of_order(client, admin_id, tmp_path):
    """Regions filters off Customers' sample (not the driver), and the links
    are given in reverse dependency order to exercise the resolver."""
    p = tmp_path / "chain.xlsx"
    _chain_workbook(p)
    ds = (await upload_file(client, admin_id, p))["dataset_id"]
    h = auth(admin_id)

    r = await client.post("/api/v1/sample/coordinated", headers=h, json=_coord_body(
        ds, target_total_volume=3,
        sampling_steps=[{"method": "random", "sample_size": 3}],
        related=[
            # listed before its parent is available — resolver must defer it
            {"sheet": "Regions", "left_on": "region_id", "right_on": "region_id",
             "parent_sheet": "Customers"},
            {"sheet": "Customers", "left_on": "customer_id",
             "right_on": "customer_id"},
        ]))
    assert r.status_code == 200, r.text
    body = r.json()
    by_sheet = {rel["sheet"]: rel for rel in body["related"]}

    driver_keys = {row["customer_id"] for row in body["driver"]["data"]}
    cust_rows = by_sheet["Customers"]["data"]
    assert {row["customer_id"] for row in cust_rows} == driver_keys

    regions = by_sheet["Regions"]
    assert regions["parent_sheet"] == "Customers"
    assert ({row["region_id"] for row in regions["data"]}
            == {row["region_id"] for row in cust_rows})


async def test_seed_makes_the_whole_slice_reproducible(client, admin_id, tmp_path):
    p = tmp_path / "chain.xlsx"
    _chain_workbook(p)
    ds = (await upload_file(client, admin_id, p))["dataset_id"]
    h = auth(admin_id)
    body = _coord_body(ds, seed=7)

    runs = []
    for _ in range(2):
        r = await client.post("/api/v1/sample/coordinated", headers=h, json=body)
        assert r.status_code == 200, r.text
        j = r.json()
        runs.append((sorted(row["order_id"] for row in j["driver"]["data"]),
                     sorted(row["customer_id"] for row in j["related"][0]["data"])))
    assert runs[0] == runs[1]


async def test_sample_files_persisted_and_fetchable(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)
    r = await client.post("/api/v1/sample/coordinated", headers=h,
                          json=_coord_body(ds))
    assert r.status_code == 200, r.text
    body = r.json()

    for fname in [body["driver"]["sample_file"],
                  body["related"][0]["sample_file"]]:
        assert fname
        raw = await client.get(f"/api/v1/samples/{fname}", headers=h)
        assert raw.status_code == 200 and len(raw.content) > 0
        rows = await client.get(f"/api/v1/samples/{fname}/data", headers=h)
        assert rows.status_code == 200 and rows.json()["total_count"] >= 0

    # The call is on the audit record like any other analytics op.
    audit = (await client.get("/api/v1/audit", params={"limit": 50},
                              headers=h)).json()
    assert any(e["path"] == "/api/v1/sample/coordinated" for e in audit["items"])


async def test_blank_driver_sheet_gets_selection_contract(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    r = await client.post("/api/v1/sample/coordinated", headers=auth(admin_id),
                          json=_coord_body(ds, driver_sheet=""))
    assert r.status_code == 400, r.text
    assert r.headers["content-type"].startswith(PROBLEM)
    body = r.json()
    assert body["code"] == "sheet-selection-required"
    assert set(body["sheets"]) == {"Customers", "Orders"}


async def test_key_and_sheet_errors(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)

    # Unknown key column on the related sheet → 400 naming the sheet.
    r = await client.post("/api/v1/sample/coordinated", headers=h, json=_coord_body(
        ds, related=[{"sheet": "Customers", "left_on": "customer_id",
                      "right_on": "nope"}]))
    assert r.status_code == 400 and "Customers" in r.json()["detail"]

    # Unknown key column on the parent (driver) side → 400 too.
    r = await client.post("/api/v1/sample/coordinated", headers=h, json=_coord_body(
        ds, related=[{"sheet": "Customers", "left_on": "nope",
                      "right_on": "customer_id"}]))
    assert r.status_code == 400 and "nope" in r.json()["detail"]

    # Unknown related sheet → 404 (sheet resolution).
    r = await client.post("/api/v1/sample/coordinated", headers=h, json=_coord_body(
        ds, related=[{"sheet": "Ghost", "left_on": "customer_id",
                      "right_on": "customer_id"}]))
    assert r.status_code == 404

    # Unknown parent_sheet → unresolvable-parents 400 (cycle/unknown guard).
    r = await client.post("/api/v1/sample/coordinated", headers=h, json=_coord_body(
        ds, related=[{"sheet": "Customers", "left_on": "customer_id",
                      "right_on": "customer_id", "parent_sheet": "Ghost"}]))
    assert r.status_code == 400 and "Unresolvable" in r.json()["detail"]


async def test_rbac_read_suffices_and_cross_team_hides(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)

    # A viewer in the owning (Default) team can run it — it's a read op.
    viewer, _ = await create_team_user(
        client, admin_id, "viewer",
        team_id="00000000-0000-0000-0000-000000000001")
    r = await client.post("/api/v1/sample/coordinated", headers=auth(viewer),
                          json=_coord_body(ds))
    assert r.status_code == 200, r.text

    # An outsider gets 404 — existence is hidden, not just forbidden.
    outsider, _ = await create_team_user(client, admin_id, "editor")
    r = await client.post("/api/v1/sample/coordinated", headers=auth(outsider),
                          json=_coord_body(ds))
    assert r.status_code == 404


# --- auto-key defaults from FK rules (ROADMAP §5) -----------------------------

async def _add_fk_rule(client, h, ds, *, name="orders-fk", sheet="Orders",
                       column="customer_id", ref_sheet="Customers",
                       ref_column="customer_id"):
    r = await client.post(f"/api/v1/datasets/{ds}/rules", headers=h, json={
        "name": name, "scope_type": "cross_sheet", "rule_type": "foreign_key",
        "sheet_selector": sheet, "column_selector": column,
        "parameters": {"ref_sheet": ref_sheet, "ref_column": ref_column},
    })
    assert r.status_code == 201, r.text
    return r.json()


async def test_omitted_keys_default_from_fk_rule(client, admin_id, tmp_path):
    """A related link without keys borrows them from the enabled FK rule."""
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)
    await _add_fk_rule(client, h, ds)

    r = await client.post("/api/v1/sample/coordinated", headers=h,
                          json=_coord_body(ds, related=[{"sheet": "Customers"}]))
    assert r.status_code == 200, r.text
    body = r.json()
    rel = body["related"][0]
    assert rel["left_on"] == "customer_id" and rel["right_on"] == "customer_id"
    driver_keys = {row["customer_id"] for row in body["driver"]["data"]}
    assert {row["customer_id"] for row in rel["data"]} == driver_keys


async def test_omitted_keys_without_rule_400(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    r = await client.post("/api/v1/sample/coordinated", headers=auth(admin_id),
                          json=_coord_body(ds, related=[{"sheet": "Customers"}]))
    assert r.status_code == 400, r.text
    assert "No foreign_key quality rule" in r.json()["detail"]


async def test_omitted_keys_ambiguous_rules_400(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)
    await _add_fk_rule(client, h, ds, name="fk-child")
    await _add_fk_rule(client, h, ds, name="fk-parent", sheet="Customers",
                       ref_sheet="Orders")
    r = await client.post("/api/v1/sample/coordinated", headers=h,
                          json=_coord_body(ds, related=[{"sheet": "Customers"}]))
    assert r.status_code == 400, r.text
    assert "Multiple foreign_key rules" in r.json()["detail"]


async def test_half_specified_keys_rejected(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    r = await client.post(
        "/api/v1/sample/coordinated", headers=auth(admin_id),
        json=_coord_body(ds, related=[{"sheet": "Customers",
                                       "left_on": "customer_id"}]))
    assert r.status_code == 422


# ---------------------------------------------------------------------------
# §24 — coordinated sampling v2: relationship-driven keys and sub-sampling
# ---------------------------------------------------------------------------

async def _confirmed_relationship(client, admin_id, ds, *, from_sheet="Orders",
                                  from_column="customer_id",
                                  to_sheet="Customers", to_column="customer_id"):
    r = await client.post(f"/api/v1/datasets/{ds}/relationships", headers=auth(admin_id),
                          json={"from_sheet": from_sheet, "from_column": from_column,
                                "to_sheet": to_sheet, "to_column": to_column})
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "confirmed"
    return r.json()


async def test_relationship_driven_keys_match_explicit_keys(client, admin_id, tmp_path):
    """A confirmed relationship resolves to exactly the keys you would type."""
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)
    rel = await _confirmed_relationship(client, admin_id, ds)

    explicit = (await client.post("/api/v1/sample/coordinated", headers=h,
                                  json=_coord_body(ds))).json()
    driven = (await client.post(
        "/api/v1/sample/coordinated", headers=h,
        json=_coord_body(ds, related=[{"sheet": "Customers",
                                       "relationship_id": rel["id"]}]))).json()

    a, b = explicit["related"][0], driven["related"][0]
    assert (b["left_on"], b["right_on"]) == (a["left_on"], a["right_on"])
    assert b["sampled_count"] == a["sampled_count"]
    # The response says where the keys came from.
    assert a["key_source"] == "explicit" and a["relationship_id"] is None
    assert b["key_source"] == "relationship" and b["relationship_id"] == rel["id"]


async def test_explicit_keys_win_over_a_relationship(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)
    rel = await _confirmed_relationship(client, admin_id, ds)
    r = await client.post("/api/v1/sample/coordinated", headers=h,
                          json=_coord_body(ds, related=[{
                              "sheet": "Customers", "relationship_id": rel["id"],
                              "left_on": "customer_id", "right_on": "customer_id"}]))
    assert r.status_code == 200, r.text
    assert r.json()["related"][0]["key_source"] == "explicit"


async def test_an_unconfirmed_relationship_cannot_drive_sampling(
        client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)
    r = await client.post(f"/api/v1/datasets/{ds}/relationships", headers=h,
                          json={"from_sheet": "Orders", "from_column": "customer_id",
                                "to_sheet": "Customers", "to_column": "customer_id",
                                "confirmed": False})
    rel = r.json()

    r = await client.post("/api/v1/sample/coordinated", headers=h,
                          json=_coord_body(ds, related=[{
                              "sheet": "Customers", "relationship_id": rel["id"]}]))
    assert r.status_code == 400
    assert r.json()["code"] == "relationship-not-confirmed"


async def test_a_relationship_not_touching_the_sheet_is_rejected(
        client, admin_id, tmp_path):
    p = tmp_path / "chain.xlsx"
    _chain_workbook(p)
    ds = (await upload_file(client, admin_id, p))["dataset_id"]
    h = auth(admin_id)
    # Regions ← Customers, which says nothing about Orders.
    rel = await _confirmed_relationship(
        client, admin_id, ds, from_sheet="Customers", from_column="region_id",
        to_sheet="Regions", to_column="region_id")

    r = await client.post("/api/v1/sample/coordinated", headers=h,
                          json={"dataset_id": ds, "driver_sheet": "Regions",
                                "target_total_volume": 2,
                                "sampling_steps": [{"method": "random", "sample_size": 2}],
                                "related": [{"sheet": "Orders",
                                             "relationship_id": rel["id"]}]})
    assert r.status_code == 400
    assert r.json()["code"] == "relationship-endpoint-mismatch"


async def test_sub_sampling_reduces_a_child_while_references_still_hold(
        client, admin_id, tmp_path):
    """§24: the child is sampled AFTER the key filter, so every surviving row
    still points at a sampled parent."""
    p = tmp_path / "chain.xlsx"
    _chain_workbook(p)
    ds = (await upload_file(client, admin_id, p))["dataset_id"]
    h = auth(admin_id)

    body = {
        "dataset_id": ds, "driver_sheet": "Customers", "target_total_volume": 5,
        "sampling_steps": [{"method": "random", "sample_size": 5}],
        "seed": 7,
        "related": [{"sheet": "Orders", "left_on": "customer_id",
                     "right_on": "customer_id",
                     "sampling_steps": [{"method": "random", "sample_size": 2}],
                     "target_total_volume": 2}],
    }
    r = await client.post("/api/v1/sample/coordinated", headers=h, json=body)
    assert r.status_code == 200, r.text
    orders = r.json()["related"][0]

    # All 5 orders are referenced by the (full) customer sample, then sub-sampled.
    assert orders["original_count"] == 5
    assert orders["referenced_count"] == 5
    assert orders["sampled_count"] == 2

    driver_ids = {row["customer_id"] for row in r.json()["driver"]["data"]}
    assert {row["customer_id"] for row in orders["data"]} <= driver_ids


async def test_without_sub_sampling_every_referenced_row_is_kept(client, admin_id, tmp_path):
    """v1 behaviour is unchanged: referenced_count == sampled_count."""
    ds = await _upload_orders(client, admin_id, tmp_path)
    h = auth(admin_id)
    r = await client.post("/api/v1/sample/coordinated", headers=h, json=_coord_body(ds))
    related = r.json()["related"][0]
    assert related["referenced_count"] == related["sampled_count"]
    assert related["key_source"] == "explicit"


async def test_sub_sampling_a_sheet_other_links_depend_on_is_rejected(
        client, admin_id, tmp_path):
    """Sub-sampling a parent would silently orphan its children."""
    p = tmp_path / "chain.xlsx"
    _chain_workbook(p)
    ds = (await upload_file(client, admin_id, p))["dataset_id"]
    h = auth(admin_id)

    r = await client.post("/api/v1/sample/coordinated", headers=h, json={
        "dataset_id": ds, "driver_sheet": "Regions", "target_total_volume": 3,
        "sampling_steps": [{"method": "random", "sample_size": 3}],
        "related": [
            # Customers is sub-sampled AND is the parent of Orders — refused.
            {"sheet": "Customers", "left_on": "region_id", "right_on": "region_id",
             "sampling_steps": [{"method": "random", "sample_size": 1}],
             "target_total_volume": 1},
            {"sheet": "Orders", "parent_sheet": "Customers",
             "left_on": "customer_id", "right_on": "customer_id"},
        ]})
    assert r.status_code == 400, r.text
    assert r.json()["code"] == "cannot-subsample-parent"
    assert r.json()["sheet"] == "Customers"


async def test_target_volume_without_steps_is_rejected(client, admin_id, tmp_path):
    ds = await _upload_orders(client, admin_id, tmp_path)
    r = await client.post("/api/v1/sample/coordinated", headers=auth(admin_id),
                          json=_coord_body(ds, related=[{
                              "sheet": "Customers", "left_on": "customer_id",
                              "right_on": "customer_id", "target_total_volume": 2}]))
    assert r.status_code == 422


async def test_sub_sampling_is_reproducible_for_a_fixed_seed(client, admin_id, tmp_path):
    p = tmp_path / "chain.xlsx"
    _chain_workbook(p)
    ds = (await upload_file(client, admin_id, p))["dataset_id"]
    h = auth(admin_id)
    body = {
        "dataset_id": ds, "driver_sheet": "Customers", "target_total_volume": 4,
        "sampling_steps": [{"method": "random", "sample_size": 4}],
        "seed": 99,
        "related": [{"sheet": "Orders", "left_on": "customer_id",
                     "right_on": "customer_id",
                     "sampling_steps": [{"method": "random", "sample_size": 2}],
                     "target_total_volume": 2}],
    }
    first = (await client.post("/api/v1/sample/coordinated", headers=h, json=body)).json()
    second = (await client.post("/api/v1/sample/coordinated", headers=h, json=body)).json()
    assert first["related"][0]["data"] == second["related"][0]["data"]
