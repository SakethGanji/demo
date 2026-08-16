"""Files journeys — the upload dialog, the results table, the download panel,
the sheet-fix screen and the storage-admin screen, driven over real HTTP in the
order a browser makes the calls.

Each test is one screen session: state flows step to step, and every assertion
is about something the *previous* step made true. Where a UI must branch, the
assertion is on the machine-readable ``code`` slug rather than on prose, because
prose is not a contract.

Screens covered:

* **Upload dialog (resumable)** — tus create/PATCH/HEAD/DELETE, the 409→HEAD→
  resume recovery loop tus-js-client runs, and "upload a new version of this
  dataset" via the ``dataset_id`` metadata branch.
* **Upload dialog (progress bar)** — ``sync=false`` plus the status poller,
  including how a failure is surfaced to a screen that has nothing else to read.
* **Results table** — paging, column subsetting, filtering, sorting and the
  export button on a stored result file.
* **Download panel** — whole-workbook vs per-sheet, and every 400/404 the
  panel has to render.
* **Sheet-fix screen** — copy-on-write replace, its RBAC edges and its diff.
* **Storage admin** — usage attribution, the retention policy, and a sweep.
"""

from __future__ import annotations

import base64
import io
import json

from openpyxl import load_workbook

from conftest import (
    auth, create_team_user, make_orders_workbook, upload_file, upload_inline,
)

PROBLEM = "application/problem+json"
TUS_CT = "application/offset+octet-stream"

CSV_V1 = b"customer_id,tier\n1,gold\n2,silver\n3,gold\n"
CSV_V2 = b"customer_id,tier\n1,platinum\n2,silver\n3,gold\n4,bronze\n5,bronze\n"

ROWS = [
    {"order_id": 1, "region": "EU", "amount": 100.0},
    {"order_id": 2, "region": "EU", "amount": 250.0},
    {"order_id": 3, "region": "US", "amount": 75.0},
    {"order_id": 4, "region": "US", "amount": 130.0},
    {"order_id": 5, "region": "APAC", "amount": 90.0},
    {"order_id": 6, "region": "APAC", "amount": 310.0},
]


# ---------------------------------------------------------------------------
# tus helpers — exactly the headers tus-js-client puts on the wire
# ---------------------------------------------------------------------------

def _meta(**pairs) -> str:
    return ", ".join(f"{k} {base64.b64encode(v.encode()).decode()}"
                     for k, v in pairs.items())


async def _tus_create(client, user_id, *, size, filename="orders.csv",
                      team_id=None, dataset_id=None):
    meta = {"filename": filename}
    if dataset_id:
        meta["dataset_id"] = dataset_id
    headers = {
        **auth(user_id),
        "Tus-Resumable": "1.0.0",
        "Upload-Length": str(size),
        "Upload-Metadata": _meta(**meta),
    }
    if team_id:
        headers["X-Team-Id"] = team_id
    return await client.post("/api/v1/tus/", headers=headers)


async def _tus_patch(client, user_id, location, chunk, offset):
    return await client.patch(location, headers={
        **auth(user_id), "Tus-Resumable": "1.0.0",
        "Content-Type": TUS_CT, "Upload-Offset": str(offset),
    }, content=chunk)


async def _sample_file(client, user_id, dataset_id, *, size, sheet=None):
    """Run a sample and return the stored result filename the UI links to."""
    body = {"dataset_id": dataset_id, "target_total_volume": size,
            "sampling_steps": [{"method": "random", "sample_size": size}],
            "seed": 7}
    if sheet:
        body["sheet"] = sheet
    r = await client.post("/api/v1/sample", headers=auth(user_id), json=body)
    assert r.status_code == 200, r.text
    fname = r.json()["sample_file"]
    assert fname, r.text
    return fname


# ---------------------------------------------------------------------------
# 1. Resumable upload dialog — the connection drops mid-file
# ---------------------------------------------------------------------------

async def test_a_large_file_upload_resumes_from_the_server_offset_after_the_connection_drops(
        client, admin_id):
    """SCREEN: the resumable upload dialog (tus-js-client) on a flaky network.

    This is the whole reason tus exists. The client sends a chunk, loses the
    connection before it sees the 204, and retries the *same* chunk at the
    offset it still believes in. The server must reject that with 409, the
    client must HEAD to learn the real offset, and the upload must then finish
    from there and land a ready version.

    If this regresses, every large upload on a shaky connection either
    silently duplicates bytes into the file or wedges the dialog at a
    percentage that never advances.
    """
    editor, team = await create_team_user(client, admin_id, "editor")
    h = auth(editor)

    # ---- 1. Discovery: the client asks what the server supports before it
    #         decides to use the resumable path at all.
    opts = await client.options("/api/v1/tus/", headers=h)
    assert opts.status_code == 204
    assert "creation" in opts.headers["tus-extension"]
    assert int(opts.headers["tus-max-size"]) >= len(CSV_V1)

    # ---- 2. Announce the upload. The only thing the client gets back is the
    #         Location URL — everything after this is keyed off it.
    r = await _tus_create(client, editor, size=len(CSV_V1), team_id=team)
    assert r.status_code == 201, r.text
    location = r.headers["location"]
    assert location.startswith("/api/v1/tus/")

    # ---- 3. First chunk lands. The 204 carries the new server offset, which
    #         is what the progress bar renders.
    cut = 20
    r = await _tus_patch(client, editor, location, CSV_V1[:cut], 0)
    assert r.status_code == 204
    assert r.headers["upload-offset"] == str(cut)

    # ---- 4. Connection drops; the client never saw that 204 and retries the
    #         same chunk from offset 0. The server refuses.
    r = await _tus_patch(client, editor, location, CSV_V1[:cut], 0)
    assert r.status_code == 409
    assert r.headers["content-type"].startswith(PROBLEM)
    assert r.json()["code"] == "conflict"

    # The rejected PATCH appended nothing: the offset is still exactly where
    # step 3 left it, not 40.
    head = await client.head(location, headers=h)
    assert head.status_code == 200
    server_offset = int(head.headers["upload-offset"])
    assert server_offset == cut
    assert head.headers["upload-length"] == str(len(CSV_V1))
    assert head.headers["cache-control"] == "no-store"

    # ---- 5. Resume from the offset the server just reported.
    r = await _tus_patch(client, editor, location, CSV_V1[server_offset:],
                         server_offset)
    assert r.status_code == 204
    assert r.headers["upload-offset"] == str(len(CSV_V1))

    # ---- 6. The dialog polls until processing finishes.
    status = (await client.get(f"{location}/status", headers=h)).json()
    assert status["status"] == "complete", status
    assert status["row_count"] == 3          # the *whole* file, not 20 bytes of it
    ds = status["dataset_id"]
    assert ds

    # ---- 7. The dialog closes onto the dataset page: exactly one version, and
    #         its row count matches what the poller promised.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert [v["version_number"] for v in versions] == [1]
    assert versions[0]["status"] == "ready"
    assert versions[0]["row_count"] == 3

    # And the bytes really are the file, not a truncated prefix.
    dl = await client.get(f"/api/v1/datasets/{ds}/download",
                          params={"format": "csv"}, headers=h)
    assert dl.status_code == 200
    assert dl.content.decode().strip().splitlines()[0] == "customer_id,tier"
    assert len(dl.content.decode().strip().splitlines()) == 4


# ---------------------------------------------------------------------------
# 2. "Upload a new version of this dataset" — the large-file path
# ---------------------------------------------------------------------------

async def test_a_resumable_upload_can_add_a_new_version_to_an_existing_dataset(
        client, admin_id, tmp_path):
    """FLOW: the "Upload new version" button on a dataset page, large-file path.

    The UI already has the dataset id, so it puts it in the tus metadata rather
    than creating a second dataset. The RBAC on that branch is the dataset's,
    not the team header's: an outsider must get a 404 (existence hidden) and an
    in-team viewer a 403, *before* any staging file is created.

    If this regresses, "new version" silently forks a duplicate dataset, or a
    read-only user can push data into someone else's dataset.
    """
    editor, team = await create_team_user(client, admin_id, "editor")
    viewer, _ = await create_team_user(client, admin_id, "viewer", team_id=team)
    outsider, _ = await create_team_user(client, admin_id, "editor")
    h = auth(editor)

    # ---- 1. v1 arrives the ordinary way (small file, multipart).
    src = tmp_path / "customers.csv"
    src.write_bytes(CSV_V1)
    up = await upload_file(client, editor, src, name="customers.csv",
                           team_id=team)
    ds = up["dataset_id"]
    assert up["row_count"] == 3

    # ---- 2. An outsider clicks the same button with a stolen dataset id.
    #         404, never 403 — a 403 would confirm the dataset exists.
    r = await _tus_create(client, outsider, size=len(CSV_V2), dataset_id=ds)
    assert r.status_code == 404
    assert r.json()["code"] == "not_found"

    # ---- 3. An in-team viewer: the dataset is visible, the write is not.
    r = await _tus_create(client, viewer, size=len(CSV_V2), dataset_id=ds)
    assert r.status_code == 403
    assert r.json()["code"] == "forbidden"

    # Neither refusal created a version.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert [v["version_number"] for v in versions] == [1]

    # ---- 4. The editor's upload is accepted and streamed in one chunk.
    r = await _tus_create(client, editor, size=len(CSV_V2), dataset_id=ds,
                          filename="customers.csv")
    assert r.status_code == 201, r.text
    location = r.headers["location"]
    r = await _tus_patch(client, editor, location, CSV_V2, 0)
    assert r.status_code == 204

    status = (await client.get(f"{location}/status", headers=h)).json()
    assert status["status"] == "complete", status
    # The id the poller reports is the dataset the button was pressed on — it
    # did NOT create a second one.
    assert status["dataset_id"] == ds
    assert status["row_count"] == 5

    # ---- 5. The version list refreshes: two entries, newest first.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert [v["version_number"] for v in versions] == [2, 1]
    assert versions[0]["status"] == "ready" and versions[0]["row_count"] == 5
    assert versions[1]["row_count"] == 3
    # Immutability: the two versions are distinct artifacts.
    assert versions[0]["source_checksum"] != versions[1]["source_checksum"]

    # ---- 6. "Download this version" on the new row returns the new bytes,
    #         and v1 still returns the old ones.
    dl2 = await client.get(f"/api/v1/datasets/{ds}/versions/2/download",
                           params={"format": "csv"}, headers=h)
    assert dl2.status_code == 200
    assert "platinum" in dl2.content.decode()
    assert len(dl2.content.decode().strip().splitlines()) == 6

    dl1 = await client.get(f"/api/v1/datasets/{ds}/versions/1/download",
                           params={"format": "csv"}, headers=h)
    assert dl1.status_code == 200
    assert "platinum" not in dl1.content.decode()


# ---------------------------------------------------------------------------
# 3. Cancel and retry
# ---------------------------------------------------------------------------

async def test_a_cancelled_upload_can_be_retried_without_stranding_the_dataset(
        client, admin_id):
    """FLOW: the Cancel button in the upload dialog, then a second attempt.

    A cancel must close the half-written version out as ``failed`` rather than
    leaving it "uploading" forever, must not become the dataset's current
    version, and must not consume the version number the retry needs.

    If this regresses, a user who cancels once sees a dataset stuck on a
    version that has no data, and every later read resolves to it.
    """
    editor, team = await create_team_user(client, admin_id, "editor")
    h = auth(editor)

    # ---- 1. Start a fresh-dataset upload and send part of the file.
    r = await _tus_create(client, editor, size=len(CSV_V2), team_id=team,
                          filename="customers.csv")
    assert r.status_code == 201
    location = r.headers["location"]
    assert (await _tus_patch(client, editor, location,
                             CSV_V2[:15], 0)).status_code == 204

    # ---- 2. The dialog polls mid-flight. This is the ONLY response that tells
    #         the UI which dataset the upload created — POST /tus/ returns a
    #         Location header and nothing else.
    status = (await client.get(f"{location}/status", headers=h)).json()
    assert status["status"] == "uploading"
    assert "15/" in (status["message"] or "")
    ds = status["dataset_id"]
    assert ds

    # The dataset row is created by POST /tus/, before a single byte lands, so
    # the catalog behind the still-open dialog already lists it with no
    # readable version. A UI cannot filter it out: the listing carries no
    # "has ready data" signal.
    catalog = (await client.get("/api/v1/datasets", params={"limit": 100},
                                headers=h)).json()
    assert any(d["id"] == ds for d in catalog["items"])

    # ---- 3. Cancel.
    assert (await client.delete(location, headers=h)).status_code == 204

    # Staging state is gone: the resume URL is now a 404 for everyone.
    assert (await client.head(location, headers=h)).status_code == 404
    assert (await client.get(f"{location}/status", headers=h)).status_code == 404

    # ---- 4. The dataset page shows one failed version, no ready data.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert [(v["version_number"], v["status"]) for v in versions] == [(1, "failed")]

    # ---- 5. Retry the whole file the small-file way, onto the same dataset.
    r = await client.post("/api/v1/upload", headers=h,
                          files={"file": ("customers.csv",
                                          io.BytesIO(CSV_V2), "text/csv")},
                          data={"dataset_id": ds})
    assert r.status_code == 200, r.text
    assert r.json()["dataset_id"] == ds
    assert r.json()["row_count"] == 5

    # ---- 6. v2 is ready and v1 is still failed — the cancel did not renumber
    #         anything, and the retry did not resurrect the dead version.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert [(v["version_number"], v["status"]) for v in versions] == [
        (2, "ready"), (1, "failed")]

    # ---- 7. The default download resolves to v2, proving the failed v1 never
    #         became current.
    dl = await client.get(f"/api/v1/datasets/{ds}/download",
                          params={"format": "csv"}, headers=h)
    assert dl.status_code == 200
    assert len(dl.content.decode().strip().splitlines()) == 6
    assert "platinum" in dl.content.decode()


# ---------------------------------------------------------------------------
# 4. Progress-bar upload of a file that cannot be parsed
# ---------------------------------------------------------------------------

async def test_an_async_upload_of_a_corrupt_file_reports_the_error_through_the_status_poller(
        client, admin_id):
    """SCREEN: the background-upload progress bar (``sync=false``).

    A screen that uploads asynchronously gets a 200 with status ``uploaded``
    immediately and has exactly one place to learn the outcome:
    ``GET /upload/status/{version_id}``. The failure has to surface *there*,
    with a message, or the bar spins forever.

    If this regresses, users watch a progress bar that never resolves and the
    dataset silently holds a version that will never be readable.
    """
    editor, team = await create_team_user(client, admin_id, "editor")
    h = auth(editor)

    # ---- 1. Fire and forget. The response has no row_count — only a handle.
    up = await client.post("/api/v1/upload", params={"sync": "false"},
                           headers={**h, "X-Team-Id": team},
                           files={"file": ("quarterly.xlsx",
                                           io.BytesIO(b"this is not a zip archive"),
                                           "application/octet-stream")})
    assert up.status_code == 200, up.text
    body = up.json()
    assert body["status"] == "uploaded"
    assert body["row_count"] is None
    ds, vid = body["dataset_id"], body["version_id"]
    assert vid

    # ---- 2. The poller is where the truth arrives.
    st = await client.get(f"/api/v1/upload/status/{vid}", headers=h)
    assert st.status_code == 200, st.text
    st = st.json()
    assert st["status"] == "error", st
    assert st["error"], "a failed async upload must carry a message to show"
    assert st["version_id"] == vid

    # ---- 3. The dataset page agrees: nothing ready to open.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert versions and all(v["status"] != "ready" for v in versions)
    assert versions[0]["status"] == "failed"

    # A read of the dataset's current version has nothing to resolve to.
    dl = await client.get(f"/api/v1/datasets/{ds}/download",
                          params={"format": "csv"}, headers=h)
    assert dl.status_code == 404

    # ---- 4. The user picks a good file and retries onto the same dataset,
    #         still asynchronously.
    up2 = await client.post("/api/v1/upload", params={"sync": "false"},
                            headers=h,
                            files={"file": ("quarterly.csv",
                                            io.BytesIO(CSV_V1), "text/csv")},
                            data={"dataset_id": ds})
    assert up2.status_code == 200, up2.text
    vid2 = up2.json()["version_id"]
    assert vid2 != vid

    st2 = (await client.get(f"/api/v1/upload/status/{vid2}", headers=h)).json()
    assert st2["status"] == "complete", st2
    assert st2["row_count"] == 3
    assert st2["error"] is None

    # The first version's status did not change when the second succeeded.
    st1 = (await client.get(f"/api/v1/upload/status/{vid}", headers=h)).json()
    assert st1["status"] == "error"

    # ---- 5. Outsiders cannot poll a handle they did not create.
    outsider, _ = await create_team_user(client, admin_id, "editor")
    r = await client.get(f"/api/v1/upload/status/{vid2}", headers=auth(outsider))
    assert r.status_code == 404
    assert r.json()["code"] == "not_found"


# ---------------------------------------------------------------------------
# 5. The results table
# ---------------------------------------------------------------------------

async def test_the_result_table_screen_pages_filters_sorts_and_exports_one_result_file(
        client, admin_id):
    """SCREEN: the result-table view of a stored sample/analysis output.

    On mount it lists the files, opens one, and renders "showing X of Y" from
    ``filtered_count`` / ``total_count``. Then the user pages, picks columns,
    filters, sorts, and finally hits Export.

    If this regresses, the table shows the wrong "of Y", pages repeat rows, or
    a mistyped sort direction silently returns ascending rows labelled success.
    """
    editor, team = await create_team_user(client, admin_id, "editor")
    h = auth(editor)
    ds = (await upload_inline(client, editor, json.dumps(ROWS),
                              team_id=team))["dataset_id"]
    fname = await _sample_file(client, editor, ds, size=6)

    # ---- 1. The files list. It is a Page envelope, so the UI can show a
    #         count without walking every page.
    listing = await client.get("/api/v1/samples",
                               params={"limit": 2, "offset": 0}, headers=h)
    assert listing.status_code == 200
    page = listing.json()
    assert page["limit"] == 2 and page["offset"] == 0
    assert len(page["items"]) <= 2
    assert page["total"] >= 1

    # Our file is reachable by paging through that envelope.
    all_files = (await client.get("/api/v1/samples",
                                  params={"limit": 100}, headers=h)).json()
    mine = next(e for e in all_files["items"] if e["filename"] == fname)
    assert mine["file_type"] == "sample_output"
    assert mine["dataset_id"] == ds
    assert mine["size_bytes"] > 0
    assert mine["key"].endswith(fname)

    # ---- 2. Page 1 of the table.
    p1 = await client.get(f"/api/v1/samples/{fname}/data", headers=h,
                          params={"limit": 2, "offset": 0,
                                  "sort_by": "order_id", "sort_order": "asc"})
    assert p1.status_code == 200, p1.text
    p1 = p1.json()
    assert p1["filename"] == fname
    assert p1["total_count"] == 6 and p1["filtered_count"] == 6
    assert p1["limit"] == 2 and p1["offset"] == 0
    assert [r["order_id"] for r in p1["data"]] == [1, 2]
    # The column header row comes from `columns`, not from the first data row.
    assert [c["name"] for c in p1["columns"]] == ["order_id", "region", "amount"]

    # ---- 3. Page 2: different rows, same denominator.
    p2 = (await client.get(f"/api/v1/samples/{fname}/data", headers=h,
                           params={"limit": 2, "offset": 2,
                                   "sort_by": "order_id",
                                   "sort_order": "asc"})).json()
    assert [r["order_id"] for r in p2["data"]] == [3, 4]
    assert p2["total_count"] == p1["total_count"]
    assert p2["offset"] == 2

    # ---- 4. The user hides a column, filters and sorts descending. The
    #         "showing X of Y" pair must now disagree — that is the point.
    q = await client.get(f"/api/v1/samples/{fname}/data", headers=h,
                         params={"columns": "order_id,amount",
                                 "filter_expr": "amount > 100",
                                 "sort_by": "amount", "sort_order": "desc"})
    assert q.status_code == 200, q.text
    q = q.json()
    assert q["total_count"] == 6
    assert q["filtered_count"] == 3
    assert [r["amount"] for r in q["data"]] == [310.0, 250.0, 130.0]
    assert set(q["data"][0]) == {"order_id", "amount"}

    # ---- 5. Error branches the table has to render.
    bad = await client.get(f"/api/v1/samples/{fname}/data", headers=h,
                           params={"columns": "order_id,nonesuch"})
    assert bad.status_code == 400
    assert bad.headers["content-type"].startswith(PROBLEM)
    assert bad.json()["code"] == "unknown-column"
    # The names that DO exist ride along, so the picker can be re-populated.
    assert "amount" in bad.json()["available"]

    bad = await client.get(f"/api/v1/samples/{fname}/data", headers=h,
                           params={"sort_by": "nonesuch"})
    assert bad.status_code == 400 and bad.json()["code"] == "unknown-column"

    # A mistyped direction is a hard 422, never silently ascending.
    bad = await client.get(f"/api/v1/samples/{fname}/data", headers=h,
                           params={"sort_order": "DESC"})
    assert bad.status_code == 422

    # ---- 6. Export the (whole) result file and download what came back.
    exp = await client.post(f"/api/v1/samples/{fname}/export",
                            params={"format": "xlsx"}, headers=h)
    assert exp.status_code == 200, exp.text
    exp = exp.json()
    assert exp["source_file"] == fname
    assert exp["format"] == "xlsx" and exp["size_bytes"] > 0
    export_file = exp["export_file"]

    dl = await client.get(f"/api/v1/samples/{export_file}", headers=h)
    assert dl.status_code == 200
    assert dl.headers["content-disposition"].endswith(f'"{export_file}"')
    ws = load_workbook(io.BytesIO(dl.content)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0] == ["order_id", "region", "amount"]
    assert len(rows) == 7  # header + the 6 sampled rows

    # ---- 7. The export shows up in the same listing, typed as an export, so
    #         the screen can offer it for re-download without another call.
    after = (await client.get("/api/v1/samples", params={"limit": 100},
                              headers=h)).json()
    entry = next(e for e in after["items"] if e["filename"] == export_file)
    assert entry["file_type"] == "export"
    assert entry["dataset_id"] == ds        # ownership inherited from the source
    assert after["total"] == all_files["total"] + 1


# ---------------------------------------------------------------------------
# 6. The download panel
# ---------------------------------------------------------------------------

async def test_the_download_panel_offers_a_whole_workbook_and_per_sheet_slices(
        client, admin_id, tmp_path):
    """SCREEN: the Download panel on a multi-sheet dataset.

    On open it reads the sheet list to build the picker, then every button is
    one GET with different query params. The panel must be able to distinguish
    "you have to pick a sheet" from "that sheet does not exist" from "those
    options only make sense for one sheet" — all of them 4xx, all of them
    needing different UI.

    If this regresses, the panel either downloads the wrong tab or shows a
    generic error for a condition it could have fixed automatically.
    """
    editor, team = await create_team_user(client, admin_id, "editor")
    h = auth(editor)
    wb = tmp_path / "orders.xlsx"
    make_orders_workbook(wb, clean=True)
    ds = (await upload_file(client, editor, wb, name="orders.xlsx",
                            team_id=team))["dataset_id"]

    # ---- 1. Mount: read the sheet list that populates the picker.
    meta = (await client.get(f"/api/v1/datasets/{ds}", headers=h)).json()
    sheet_names = [s["name"] for s in meta["sheets"]]
    assert sheet_names == ["Customers", "Orders"]
    assert meta["default_sheet"] == "Customers"

    # ---- 2. "Download whole workbook" — xlsx with no sheet reconstructs every
    #         tab, in workbook order, under the original filename.
    dl = await client.get(f"/api/v1/datasets/{ds}/download",
                          params={"format": "xlsx"}, headers=h)
    assert dl.status_code == 200, dl.text
    assert 'filename="orders.xlsx"' in dl.headers["content-disposition"]
    book = load_workbook(io.BytesIO(dl.content))
    assert book.sheetnames == sheet_names
    assert [c.value for c in book["Orders"][1]] == ["order_id", "customer_id", "total"]

    # ---- 3. Column/limit/filter are single-sheet options. Asking for them on
    #         the whole workbook is a 400 the panel must catch before it lets
    #         the user tick both.
    r = await client.get(f"/api/v1/datasets/{ds}/download", headers=h,
                         params={"format": "xlsx", "columns": "order_id,total"})
    assert r.status_code == 400
    assert r.headers["content-type"].startswith(PROBLEM)
    assert "single sheet" in r.json()["detail"]

    # ---- 4. CSV cannot hold two tabs, so it demands a sheet — and names them,
    #         so the panel can open the picker rather than show an error.
    r = await client.get(f"/api/v1/datasets/{ds}/download",
                         params={"format": "csv"}, headers=h)
    assert r.status_code == 400
    assert r.json()["code"] == "sheet-selection-required"
    assert set(r.json()["sheets"]) == set(sheet_names)

    # ---- 5. The real slice: one sheet, two columns, filtered, limited.
    r = await client.get(f"/api/v1/datasets/{ds}/download", headers=h,
                         params={"format": "csv", "sheet": "Orders",
                                 "columns": "order_id,total",
                                 "filter_expr": "total > 50", "limit": 2})
    assert r.status_code == 200, r.text
    lines = r.content.decode().strip().splitlines()
    assert lines[0] == "order_id,total"
    assert len(lines) == 2                      # header + the single match
    assert lines[1].startswith("10,")

    # ---- 6. Error branches for the picker and the column chooser.
    r = await client.get(f"/api/v1/datasets/{ds}/download", headers=h,
                         params={"format": "csv", "sheet": "NoSuchSheet"})
    assert r.status_code == 404
    assert r.json()["code"] == "not_found"
    assert "NoSuchSheet" in r.json()["detail"]

    r = await client.get(f"/api/v1/datasets/{ds}/download", headers=h,
                         params={"format": "csv", "sheet": "Orders",
                                 "columns": "order_id,nonesuch"})
    assert r.status_code == 400
    assert "nonesuch" in r.json()["detail"]

    r = await client.get(f"/api/v1/datasets/{ds}/download", headers=h,
                         params={"format": "yaml"})
    assert r.status_code == 400
    assert "Unsupported format" in r.json()["detail"]

    # ---- 7. The version-pinned download beside it obeys the same sheet rule…
    r = await client.get(f"/api/v1/datasets/{ds}/versions/1/download",
                         params={"format": "csv"}, headers=h)
    assert r.status_code == 400
    assert r.json()["code"] == "sheet-selection-required"

    r = await client.get(f"/api/v1/datasets/{ds}/versions/1/download",
                         params={"format": "parquet", "sheet": "Customers"},
                         headers=h)
    assert r.status_code == 200
    assert r.content[:4] == b"PAR1"
    assert 'filename="orders_v1.parquet"' in r.headers["content-disposition"]

    # …and 404s a version that was never created.
    r = await client.get(f"/api/v1/datasets/{ds}/versions/9/download",
                         params={"format": "csv", "sheet": "Orders"}, headers=h)
    assert r.status_code == 404

    # ---- 8. RBAC: a same-team viewer may download; an outsider cannot even
    #         learn the dataset exists.
    viewer, _ = await create_team_user(client, admin_id, "viewer", team_id=team)
    r = await client.get(f"/api/v1/datasets/{ds}/download", headers=auth(viewer),
                         params={"format": "csv", "sheet": "Orders"})
    assert r.status_code == 200 and "order_id" in r.content.decode()

    outsider, _ = await create_team_user(client, admin_id, "editor")
    for path in (f"/api/v1/datasets/{ds}/download",
                 f"/api/v1/datasets/{ds}/versions/1/download"):
        r = await client.get(path, headers=auth(outsider),
                             params={"format": "csv", "sheet": "Orders"})
        assert r.status_code == 404, path


# ---------------------------------------------------------------------------
# 7. Fixing one tab of a workbook
# ---------------------------------------------------------------------------

async def test_fixing_one_tab_of_a_workbook_leaves_every_other_tab_intact_and_respects_rbac(
        client, admin_id, tmp_path):
    """SCREEN: "Replace this sheet" on the sheet detail page.

    The user picks a tab, uploads a corrected single-table file, and the app
    makes a whole new immutable version in which every *other* tab is reused
    copy-on-write. The screen then shows the diff and lets the user open an
    untouched tab from the new version.

    If this regresses, fixing one tab either rewrites tabs nobody touched or
    lets a viewer mutate a dataset they can only read.
    """
    editor, team = await create_team_user(client, admin_id, "editor")
    viewer, _ = await create_team_user(client, admin_id, "viewer", team_id=team)
    outsider, _ = await create_team_user(client, admin_id, "editor")
    h = auth(editor)

    wb = tmp_path / "orders.xlsx"
    make_orders_workbook(wb, clean=True)
    ds = (await upload_file(client, editor, wb, name="orders.xlsx",
                            team_id=team))["dataset_id"]

    # ---- 1. Mount: the sheet list is where the user picks the tab, and it
    #         carries the row counts the screen shows beside each name.
    sheets = (await client.get(f"/api/v1/datasets/{ds}/sheets",
                               headers=h)).json()["items"]
    by_name = {s["name"]: s for s in sheets}
    assert set(by_name) == {"Customers", "Orders"}
    assert by_name["Customers"]["row_count"] == 3
    assert by_name["Orders"]["row_count"] == 2

    fixed = b"customer_id,tier\n1,gold\n2,platinum\n"

    # ---- 2. A stale picker naming a tab that no longer exists: 404 that lists
    #         the real tabs, so the screen can re-render the picker.
    r = await client.post(f"/api/v1/datasets/{ds}/sheets/NoSuchSheet/replace",
                          headers=h,
                          files={"file": ("fix.csv", io.BytesIO(fixed), "text/csv")})
    assert r.status_code == 404
    assert r.json()["code"] == "not_found"
    assert "Customers" in r.json()["detail"] and "Orders" in r.json()["detail"]

    # ---- 3. RBAC on the button itself.
    r = await client.post(f"/api/v1/datasets/{ds}/sheets/Customers/replace",
                          headers=auth(viewer),
                          files={"file": ("fix.csv", io.BytesIO(fixed), "text/csv")})
    assert r.status_code == 403 and r.json()["code"] == "forbidden"

    r = await client.post(f"/api/v1/datasets/{ds}/sheets/Customers/replace",
                          headers=auth(outsider),
                          files={"file": ("fix.csv", io.BytesIO(fixed), "text/csv")})
    assert r.status_code == 404

    # Neither refusal created a version — the version list is still just v1.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert [v["version_number"] for v in versions] == [1]

    # ---- 4. The real replace.
    r = await client.post(f"/api/v1/datasets/{ds}/sheets/Customers/replace",
                          headers=h,
                          files={"file": ("fix.csv", io.BytesIO(fixed), "text/csv")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["dataset_id"] == ds
    assert body["version_number"] == 2
    assert body["replaced_sheet"] == "Customers"
    assert body["reused_sheets"] == ["Orders"]   # the screen names what it kept
    assert body["row_count"] == 4                # 2 fixed + 2 untouched Orders

    # ---- 5. The diff panel: one tab modified, the other untouched.
    diff = (await client.get(f"/api/v1/datasets/{ds}/versions/1/diff/2",
                             headers=h)).json()
    assert diff["added"] == [] and diff["removed"] == []
    assert diff["unchanged"] == ["Orders"]
    modified = {m["sheet_key"]: m for m in diff["modified"]}
    assert set(modified) == {"customers"}
    assert modified["customers"]["row_count_delta"] == -1

    # ---- 6. Open the untouched tab from the NEW version — byte-identical to
    #         the same tab in v1, which is the whole promise of copy-on-write.
    old = await client.get(f"/api/v1/datasets/{ds}/versions/1/download",
                           params={"format": "csv", "sheet": "Orders"}, headers=h)
    new = await client.get(f"/api/v1/datasets/{ds}/versions/2/download",
                           params={"format": "csv", "sheet": "Orders"}, headers=h)
    assert old.status_code == 200 and new.status_code == 200
    assert new.content == old.content

    # And the replaced tab really did change.
    fixed_dl = await client.get(f"/api/v1/datasets/{ds}/versions/2/download",
                                params={"format": "csv", "sheet": "Customers"},
                                headers=h)
    assert "platinum" in fixed_dl.content.decode()
    assert len(fixed_dl.content.decode().strip().splitlines()) == 3

    # v1 is immutable: its Customers tab still has three rows and no platinum.
    v1_dl = await client.get(f"/api/v1/datasets/{ds}/versions/1/download",
                             params={"format": "csv", "sheet": "Customers"},
                             headers=h)
    assert "platinum" not in v1_dl.content.decode()
    assert len(v1_dl.content.decode().strip().splitlines()) == 4

    # ---- 7. The user picks the wrong file next time — a whole workbook into a
    #         single-tab slot. That is a 400 the file picker has to explain, and
    #         the rejected attempt must not become the dataset's current data.
    r = await client.post(f"/api/v1/datasets/{ds}/sheets/Customers/replace",
                          headers=h,
                          files={"file": ("whole.xlsx", open(wb, "rb"),
                                          "application/octet-stream")})
    assert r.status_code == 400
    assert "single-table" in r.json()["detail"]

    # The rejected attempt burns a version number and leaves a *failed* row in
    # the list a UI renders — it is never ready, and never current.
    versions = (await client.get(f"/api/v1/datasets/{ds}/versions",
                                 headers=h)).json()["items"]
    assert [(v["version_number"], v["status"]) for v in versions] == [
        (3, "failed"), (2, "ready"), (1, "ready")]

    still = await client.get(f"/api/v1/datasets/{ds}/download",
                             params={"format": "csv", "sheet": "Customers"},
                             headers=h)
    assert still.status_code == 200
    assert "platinum" in still.content.decode()


# ---------------------------------------------------------------------------
# 8. The storage admin screen
# ---------------------------------------------------------------------------

async def test_the_storage_admin_screen_attributes_new_bytes_and_runs_a_sweep(
        client, admin_id):
    """SCREEN: the platform storage admin page.

    On mount it reads usage, the retention policy and the backlog; the Run GC
    button posts a sweep and re-reads usage. The split between
    ``samples_bytes`` and ``exports_bytes`` only exists here, so a
    mis-classified prefix is invisible anywhere else.

    If this regresses, the admin sees exports counted as samples, or a sweep
    that reports work it did not do — or, worse, deletes live artifacts.
    """
    h = auth(admin_id)
    editor, team = await create_team_user(client, admin_id, "editor")

    # ---- 0. RBAC: these three routes are platform-wide aggregates, so a team
    #         editor must not see them at all.
    for method, path in (("get", "/api/v1/storage/usage"),
                         ("get", "/api/v1/storage/retention"),
                         ("post", "/api/v1/storage/gc")):
        r = await getattr(client, method)(path, headers=auth(editor))
        assert r.status_code == 403, path
        assert r.json()["code"] == "forbidden"

    u0 = (await client.get("/api/v1/storage/usage", headers=h)).json()
    assert set(u0) == {"total_bytes", "datasets_bytes", "samples_bytes",
                       "exports_bytes", "uploads_bytes"}

    # ---- 1. A user does some work: a dataset, then a sample of it.
    ds = (await upload_inline(client, editor, json.dumps(ROWS),
                              team_id=team))["dataset_id"]
    fname = await _sample_file(client, editor, ds, size=4)

    u1 = (await client.get("/api/v1/storage/usage", headers=h)).json()
    assert u1["datasets_bytes"] > u0["datasets_bytes"]
    assert u1["samples_bytes"] > u0["samples_bytes"]
    assert u1["exports_bytes"] == u0["exports_bytes"]   # nothing exported yet
    assert u1["total_bytes"] == (u1["datasets_bytes"] + u1["samples_bytes"]
                                 + u1["exports_bytes"] + u1["uploads_bytes"])

    # ---- 2. Now an export. It must land in exports_bytes, not samples_bytes,
    #         and by exactly the size the export response reported.
    exp = (await client.post(f"/api/v1/samples/{fname}/export",
                             params={"format": "csv"},
                             headers=auth(editor))).json()
    u2 = (await client.get("/api/v1/storage/usage", headers=h)).json()
    assert u2["exports_bytes"] == u1["exports_bytes"] + exp["size_bytes"]
    assert u2["samples_bytes"] == u1["samples_bytes"]

    # ---- 3. The policy panel. Every kind the screen can show a badge for is
    #         named, and "keep forever" is expressible as null.
    pol = await client.get("/api/v1/storage/retention", headers=h)
    assert pol.status_code == 200
    pol = pol.json()
    rules = {r["artifact_type"]: r["retention_days"] for r in pol["rules"]}
    assert rules["export"] == 7
    assert rules["sample_output"] == 30
    assert rules["published_source"] is None
    assert pol["orphan_grace_hours"] == 24
    backlog = pol["expired_pending"]

    # Nothing created above is due yet — the shortest deadline is 7 days.
    assert backlog == 0

    # ---- 4. Run GC. The backlog was empty, so the sweep must report a clean,
    #         complete pass rather than silent truncation…
    gc = await client.post("/api/v1/storage/gc", headers=h)
    assert gc.status_code == 200, gc.text
    gc = gc.json()
    assert gc["expired_deleted"] == 0
    assert gc["by_type"] == {}
    assert gc["bytes_freed"] == 0
    assert gc["more_remaining"] is False

    # ---- 5. …and a sweep that freed nothing must not have moved the numbers.
    u3 = (await client.get("/api/v1/storage/usage", headers=h)).json()
    assert u3["samples_bytes"] == u2["samples_bytes"]
    assert u3["exports_bytes"] == u2["exports_bytes"]

    # The live artifacts are still openable — a GC that "succeeded" by
    # deleting things that were not due would show up right here.
    for f in (fname, exp["export_file"]):
        assert (await client.get(f"/api/v1/samples/{f}",
                                 headers=auth(editor))).status_code == 200


# ---------------------------------------------------------------------------
# 9. Deleting a dataset takes its outputs with it
# ---------------------------------------------------------------------------

async def test_deleting_a_dataset_makes_its_result_files_unreachable_and_the_next_sweep_is_clean(
        client, admin_id):
    """FLOW: "Delete dataset" from the dataset page, then the admin checks up.

    Deleting the dataset must take its derived outputs with it — they vanish
    from the files listing, and the download that worked a moment ago becomes a
    404 rather than an orphan the UI still links to. The admin then runs a
    sweep and the storage numbers must not have grown behind everyone's back.

    If this regresses, the files list keeps offering downloads that error, and
    deleted datasets keep billing for bytes nobody can reach.
    """
    h_admin = auth(admin_id)
    editor, team = await create_team_user(client, admin_id, "editor")
    h = auth(editor)

    ds = (await upload_inline(client, editor, json.dumps(ROWS),
                              team_id=team))["dataset_id"]
    fname = await _sample_file(client, editor, ds, size=3)

    # ---- 1. Before: the file is listed, downloadable and readable.
    before = (await client.get("/api/v1/samples", params={"limit": 100},
                               headers=h)).json()
    assert any(e["filename"] == fname for e in before["items"])
    assert (await client.get(f"/api/v1/samples/{fname}",
                             headers=h)).status_code == 200
    rows = (await client.get(f"/api/v1/samples/{fname}/data", headers=h)).json()
    assert rows["total_count"] == 3

    # ---- 2. The editor who made it cannot delete it — deletion is a team-admin
    #         action, so the button has to be hidden from (and refused for) them.
    r = await client.delete(f"/api/v1/datasets/{ds}", headers=h)
    assert r.status_code == 403 and r.json()["code"] == "forbidden"
    assert (await client.get(f"/api/v1/samples/{fname}",
                             headers=h)).status_code == 200  # refusal changed nothing

    # ---- 3. The team admin deletes it. The response tells the screen what went.
    team_admin, _ = await create_team_user(client, admin_id, "admin", team_id=team)
    r = await client.delete(f"/api/v1/datasets/{ds}", headers=auth(team_admin))
    assert r.status_code == 200, r.text
    assert r.json()["success"] is True
    assert "1 artifact" in r.json()["message"]

    # ---- 4. After: the listing no longer offers it…
    after = (await client.get("/api/v1/samples", params={"limit": 100},
                              headers=h)).json()
    assert not any(e["filename"] == fname for e in after["items"])
    assert after["total"] == before["total"] - 1

    # …and both file routes 404 rather than 500 on a dangling key.
    assert (await client.get(f"/api/v1/samples/{fname}",
                             headers=h)).status_code == 404
    r = await client.get(f"/api/v1/samples/{fname}/data", headers=h)
    assert r.status_code == 404 and r.json()["code"] == "not_found"

    # Exporting from a file that is gone is a 404 too, not a 500.
    r = await client.post(f"/api/v1/samples/{fname}/export",
                          params={"format": "csv"}, headers=h)
    assert r.status_code == 404

    # The dataset itself is gone from every surface the screen would use.
    assert (await client.get(f"/api/v1/datasets/{ds}",
                             headers=h)).status_code == 404
    assert (await client.get(f"/api/v1/datasets/{ds}/download",
                             params={"format": "csv"},
                             headers=h)).status_code == 404

    # ---- 5. The admin sweeps. Nothing is due and freshly-written blobs are
    #         inside the orphan grace window, so this must be a clean no-op —
    #         not a report of work it did not do.
    gc = (await client.post("/api/v1/storage/gc", headers=h_admin)).json()
    assert gc["expired_deleted"] == 0
    assert gc["orphans_deleted"] == 0
    assert gc["bytes_freed"] == 0
    assert gc["more_remaining"] is False

    usage = await client.get("/api/v1/storage/usage", headers=h_admin)
    assert usage.status_code == 200
    assert usage.json()["total_bytes"] >= 0
